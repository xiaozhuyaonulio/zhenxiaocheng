"""Build Chinese report, review workbook, Word report and standalone figures."""
from pathlib import Path
import json,re,math
import numpy as np
import pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.font_manager import FontProperties
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.chart import BarChart, Reference
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]; OUT=ROOT/'analysis'
M=json.loads((OUT/'metrics.json').read_text())
def read(n): return pd.read_csv(OUT/(n+'.csv'))
asin=read('SQP_ASIN汇总'); qs=read('SQP_搜索词诊断'); aq=read('SQP_ASIN词明细')
port=read('广告_组合汇总'); types=read('广告_投放类型'); qtypes=read('广告_客户词类型')
pq=read('广告_组合词汇总'); targets=read('广告_投放单元汇总'); cross=read('两表_全部词对照')
queues=read('行动_待复核队列'); coverage=read('两表_集合覆盖'); checks=read('质量_核对记录')
def money(x): return f'${x:,.2f}'
def pct(x): return f'{x:.1%}'
def table(d, cols, fmts=None):
    fmts=fmts or {}
    lines=['| '+' | '.join(cols)+' |','| '+' | '.join(['---']*len(cols))+' |']
    for _,r in d.iterrows():
        vals=[]
        for c in cols:
            v=r[c]
            if pd.isna(v): v='—'
            elif c in fmts: v=fmts[c](v)
            elif isinstance(v,(float,np.floating)): v=f'{v:,.0f}' if v==int(v) else f'{v:.2f}'
            vals.append(str(v).replace('|','/').replace('\n',' '))
        lines.append('| '+' | '.join(vals)+' |')
    return '\n'.join(lines)

evidence=pd.DataFrame([
 ['已验证事实','从原始计数汇总、同词匹配和计算校验得到','花费、销售额、点击、购买计数、ACOS、日期区间、词条覆盖','可以复算，不等于因果或完整店铺规模。'],
 ['诊断线索','原始计数基础上构造描述性对照','同词购买率差异、结构调整购买指数、机会矩阵','用于排序检查；不确定原因，不表示统计显著。'],
 ['待验证建议','筛选规则与业务判断','增长候选、高成本复核候选、调价或预算测试','25%/40%只是筛选线，未提供目标ACOS与利润，不能作为盈亏或收益承诺。'],
 ['当前不能确定','缺少对齐口径或直接证据','自然流量占比、广告增量、净利润、页面因素的因果影响、计数加噪','不把脱敏等推测当成差异原因；补充官方定义和业务数据后再判断。'],
],columns=['证据层级','依据','例子','解释边界'])

notes=pd.DataFrame([
 ['分析范围','原始文件载明2026年9月；51个SQP文件3,117行、2,410词；美国USD广告61,169行。'],
 ['最重要边界','SQP为已导出ASIN×搜索词样本；广告为本份SP报告汇总。两者不能相减估自然流量。'],
 ['广告日期','开始—结束是每行覆盖区间，不能当日级发生数据；月合计保留，日趋势取消。'],
 ['市场去重','同一词市场搜索/曝光/点击/加购/购买计数完全一致；全局市场按归一化词仅计一次。'],
 ['未收录不是0','两表对照中缺失留空；集合表对应方无记录也留空；只对已出现且原始计数0的记录写0。'],
 ['归一化','仅转小写、去首尾空格、合并连续空白；保留单复数、语种、拼写差异，不做语义合并。'],
 ['广告客户词类型','完全匹配b0加8个英数字的条目标ASIN编码；编码不是普通文本搜索词，位置仍需报告确认。'],
 ['广告投放分类','先按close-match/loose-match/substitutes/complements，再按匹配类型和投放表达式判断。-不自动等于自动广告。'],
 ['广告组合≠ASIN','组合按原名称统计；本表无推广ASIN/SKU，名称内编码不作为可靠商品映射，也不假设完整店铺范围。'],
 ['SQP率的分母','原始市场点击率=市场点击/搜索量，加购率=加购/搜索量，购买率=购买/搜索量。'],
 ['自算率','样本点击率=点击/曝光；点击后购买率=购买计数/点击。另命名，避免与原始率混淆。'],
 ['结构调整购买指数','实际购买 / Σ(该ASIN每词点击×同词市场购买/市场点击)。控制观测词结构；不控制价格、位置、人群、缺失词，不是因果或显著性检验。'],
 ['相对市场购买指数','同词样本点击后购买率/同词市场点击后购买率；指数高也要看点击和购买样本数。'],
 ['购买差额','当前点击按市场率预期购买−已观察购买；仅排序诊断，不能当可挽回订单或预测收入。'],
 ['行动阈值','零单≥20点击且≥$20；有单≥5且花费≥$50、ACOS>40%；增长≥10单、ACOS≤25%。为复核筛选线，用户未提供目标ACOS。'],
 ['执行边界','输出供带教复核；先看投放单元、推广商品、库存和利润，再测试调价/预算/否定。未修改任何广告设置。'],
 ['追溯','SQP明细保留来源文件和原始CSV行号（首条数据第3行）；广告组合词保留原始Excel首末行号，首末仅索引范围并非连续匹配。'],
 ['完整明细','全44,454条两表对照、51,736条组合词明细、43,846条广告词汇总保留在同目录CSV；工作簿收录SQP相关2,410条对照及176条行动候选。'],
 ['读取原表','source_material/raw_data保存原始CSV和Excel，文件内容未改动。所有计算仅使用这些原表。'],
],columns=['主题','解释'])

# Human-review actions are specific enough to inspect in the advertising console.
actions=pd.DataFrame([
 ['P1','先查大体量承接','SQP B0DGCVK1JL；mason jars / cold brew pitcher','前者1,465点击57购买，购买率3.9% vs 同词市场12.4%；后者960点击50购买，5.2% vs12.0%','核查词意图、商品展示位置、主图、规格说明、价格优惠、评价和缺货；用后台证据确定瓶颈','同词同ASIN点击后购买率、购买份额；不把流量差当自然/广告差'],
 ['P1','复核同词不同组合成本','mason jar pitcher：梅森罐 vs32oz梅森罐','22.9%/437广告订单 vs47.3%/91订单；不能仅按词统一决策','先在投放单元和推广商品层核对；高成本组合试降出价10%—15%，低成本组合有预算约束才试增预算10%','归因成熟后的ACOS、订单/日、广告SKU和其他SKU销售构成'],
 ['P1','复核高花费同类投放','【梅森罐】低价扫流0.7 -2 - z / substitutes','花费3,654.70，346单，ACOS41.9%','拆看客户ASIN和推广商品，把差的具体商品目标列复核；避免整类一刀切','同类定向的ACOS、订单；是否损失已有高效目标'],
 ['P1','复核32oz关联投放','【32oz梅森罐】低价扫流0.7 - z / complements','花费944.69，77单，ACOS53.3%；该类总ACOS46.3%且62.7%销售来自其他SKU','检查关联目标、落地商品和连带利润；只对高成本单元试降价','总ACOS与SKU结构同时跟踪，不能把其他SKU销售全删掉判断'],
 ['P2','增长测试候选','【梅森罐】手动精准核心1 mason jar pitcher / EXACT','投放单元花费2,131.72，374单，ACOS22.3%','确认目标ACOS、库存及预算是否受限；预算受限才试增10%，出价与预算分开测试','新增花费对应订单、ACOS和总销量；不能保证边际效率不变'],
 ['P2','另一个增长候选','【梅森罐】手动广泛 glass pitchers with handle and lid - z / BROAD','投放单元花费4,029.91，633单，ACOS23.6%','保留高效广泛探索，拆出表现好的客户词；新增精准组前检查现有覆盖','客户词、精准/广泛分流后的ACOS和订单，不因为匹配类型统一迁移'],
 ['P2','复核零单词/编码','26个组合×词候选，总花费702.04','仅占全表花费0.57%；20点击零单仍可能是随机波动','文本词核相关性，ASIN编码核商品目标；按具体活动/广告组处理，避免全店否词','归因成熟后的点击/订单，是否持续低效；不宣称50,472.72全是浪费'],
 ['P2','扩展适配词','glass carafe with lid；64oz黑色B0DGCVK1JL','全样本同词145点击21购买；其中64oz黑色144点击21购买','检查实际壶型/盖子/容量是否满足carafe意图；先小额词测试和文案相关性检查','同ASIN同词转化、广告成本；145点击样本较小，先验证'],
 ['P3','补齐口径','日粒度SP；推广商品报告；SQP更多词/品牌视图；利润','缺日度、缺可靠ASIN映射、缺盈亏线，不能算自然份额或利润','申请当前权限可导出的更多SQP记录，保留官方字段说明；不要预设一定能导出1,000条','同市场同日期同ASIN范围后重做对照；仍按报告归因说明判断能否相减'],
],columns=['优先级','任务','对象','证据','建议测试或核查','观察指标及边界'])
actions.to_csv(OUT/'带教讨论行动清单.csv',index=False,encoding='utf-8-sig')

# Figures: static, exportable, with descriptive captions rather than inferred causality.
font_paths=[Path('/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc'),
    Path('C:/Windows/Fonts/msyh.ttc'),Path('/System/Library/Fonts/PingFang.ttc')]
font_path=next((f for f in font_paths if f.exists()),None)
font=FontProperties(fname=str(font_path)) if font_path else FontProperties(family='sans-serif')
plt.rcParams.update({'font.family':font.get_name(),'axes.unicode_minus':False,'figure.dpi':140,'savefig.dpi':180})
plt.rcParams['font.sans-serif']=[font.get_name()]
BLUE='#245b83'; ORANGE='#d78036'; GREEN='#3b887a'
fig,axes=plt.subplots(1,2,figsize=(12,4.2),layout='constrained')
vals=[asin.iloc[0]['样本点击'],asin['样本点击'].sum()-asin.iloc[0]['样本点击']]
axes[0].bar(['64oz黑色单装','其余50个ASIN'],vals,color=[BLUE,'#b4c9d8'])
axes[0].set_title('SQP样本点击集中度'); axes[0].set_ylabel('点击计数')
for i,v in enumerate(vals): axes[0].text(i,v,f'{v:,}\n{v/sum(vals):.1%}',ha='center',va='bottom')
sub=asin.iloc[[0,1,2,3]].copy()
labs=['64oz黑色','32oz黑色','32oz绿色','68oz冰箱壶']
xx=np.arange(len(sub)); axes[1].bar(xx-.18,sub['点击后购买率']*100,.36,label='样本购买/点击',color=BLUE)
axes[1].bar(xx+.18,sub['按自身点击词结构市场购买率']*100,.36,label='同词市场率加权',color=ORANGE)
axes[1].set_xticks(xx,labs);axes[1].set_ylabel('%');axes[1].set_title('按各ASIN自身点击词结构对照');axes[1].legend(fontsize=9)
fig.suptitle('2026年9月｜已导出SQP样本，非全店全量；率差不能解释为颜色的因果影响',fontsize=11)
fig.savefig(OUT/'图1_商品集中与承接.png');plt.close(fig)

fig,ax=plt.subplots(figsize=(9.5,4.3),layout='constrained')
labels=['两表都收录','仅广告文本词','仅广告ASIN编码']; values=[56382,55995,46191]
ax.barh(labels,values,color=[BLUE,ORANGE,GREEN]);ax.invert_yaxis();ax.set_xlim(0,68000)
for i,v in enumerate(values): ax.text(v+800,i,f'{v:,}（{v/158568:.1%}）',va='center',fontsize=10)
ax.set_xlabel('本份SP报告广告点击计数');ax.set_title('广告点击覆盖：不能把未匹配条目全部称为Top100之外的长尾词')
fig.savefig(OUT/'图2_广告点击覆盖.png');plt.close(fig)

fig,ax=plt.subplots(figsize=(9.5,4.6),layout='constrained')
names=['【梅森罐】','【果茶款梅森罐】','【滤棒款梅森罐】','【灰色梅森】','【32oz梅森罐】']
sub=pq[pq['搜索词'].eq('mason jar pitcher')].set_index('广告组合名称').loc[names]
ax.barh([x.strip('【】') for x in names],sub['ACOS']*100,color=[BLUE,ORANGE,ORANGE,ORANGE,ORANGE]);ax.invert_yaxis();ax.set_xlim(0,61)
for i,(_,r) in enumerate(sub.iterrows()): ax.text(r['ACOS']*100+.7,i,f"{r['ACOS']:.1%} / {int(r['广告订单'])}单",va='center')
ax.set_xlabel('总ACOS（%，含其他SKU销售额）');ax.set_title('同一客户词 mason jar pitcher：不同广告组合效率差异')
fig.savefig(OUT/'图3_同词不同组合.png');plt.close(fig)

fig,ax=plt.subplots(figsize=(10,5.2),layout='constrained')
sub=types[types['广告花费'].ge(800)].sort_values('ACOS');lab=sub['投放类型'].tolist()
ax.barh(lab,sub['ACOS']*100,color=[BLUE if x<.3 else ORANGE for x in sub['ACOS']]);ax.invert_yaxis();ax.set_xlim(0,60)
for i,(_,r) in enumerate(sub.iterrows()): ax.text(r['ACOS']*100+.6,i,f"{r['ACOS']:.1%} / 花费{money(r['广告花费'])}",va='center',fontsize=9)
ax.set_xlabel('总ACOS（%）；不同商品、词和位置构成影响比较');ax.set_title('投放类型复核：先区分自动、手动词和商品定向')
fig.savefig(OUT/'图4_投放类型效率.png');plt.close(fig)

fmt={'广告花费':money,'广告销售额':money,'ACOS':pct,'其他SKU销售额占比':pct,'点击后购买率':pct,'结构调整购买指数':lambda x:f'{x:.2f}倍',
     '样本点击后购买率':pct,'市场点击后购买率':pct,'购买份额':pct,'相对市场购买指数':lambda x:f'{x:.2f}倍'}
same=pq[(pq['搜索词']=='mason jar pitcher')&pq['广告组合名称'].isin(names)][['广告组合名称','广告花费','广告订单','ACOS']].sort_values('ACOS')
gaps=qs[qs['搜索词'].isin(['mason jars','cold brew pitcher','mason jar','glass water bottles','kitchen','home'])]
zero=queues[queues['建议队列'].str.startswith('优先')].head(8)
weak=queues[queues['建议队列'].str.startswith('复核')&queues['客户词类型'].eq('文本词')].head(6)
grow=queues[queues['建议队列'].str.startswith('增长')&queues['客户词类型'].eq('文本词')].head(6)

report=f'''# 2026年9月亚马逊搜索与广告分析

分析人：独立复核原始文件；供与带教 zhenxiaocheng 讨论。范围：美国站广告、51个ASIN的SQP导出样本。广告金额均为USD。

## 先看结论

本报告直接从raw_data中的51个SQP CSV、广告原表和商品对应关系重新计算，围绕商品承接、广告降本与增长测试展开。当前数据适合回答：哪些商品的搜索承接值得优先检查、哪些词与商品组合投得有效、哪些投放单元应复核成本。当前数据无法可靠回答自然流量占比、广告增量贡献或净利润。

1. **先优化主力商品的承接。** B0DGCVK1JL（64oz黑色单装梅森罐）占SQP样本70.5%的点击、61.9%的购买计数。同词结构对照后，购买效率指数为0.77，优先检查它获得的流量是否适配商品，以及页面承接是否充分。
2. **按“搜索词×广告组合×投放单元”做决策。** `mason jar pitcher`在【梅森罐】组合的ACOS为22.9%，在【32oz梅森罐】为47.3%。同词好坏取决于落地商品、投放方式和成本，不能全店统一加价或否词。
3. **降本先看有体量的低效投放。** 自动关联整体ACOS46.3%，自动同类36.7%；【梅森罐】一条同类单元花费$3,654.70、ACOS41.9%。先拆具体目标，避免按整个类型删除。
4. **总ACOS需要结合SKU销售结构。** 全表花费$122,725.96、7天归因销售额$412,829.06，总ACOS29.7%。其他SKU销售额占44.0%，花费÷广告SKU销售额为53.1%；这不是利润亏损结论，也不是把其他SKU销售额判为无价值。
5. **增长候选可以测试，不能直接预测收益。** 105个组合×词候选满足“≥10单、ACOS≤25%”；例如【梅森罐】的`pitcher`、`glass pitcher with lid`。先确定利润和库存、是否受预算限制，再小步测试。

## 1. 数据能说明什么

| 数据 | 已验证范围 | 适合回答的问题 | 主要限制 |
| --- | --- | --- | --- |
| 51个SQP CSV | 3,117行、2,410个去重词；27个文件达到100行，排名1—100 | 指定词的市场规模、样本ASIN表现、点击与购买承接 | 前100导出；不足100行也不证明全量；不能代表所有搜索或全店 |
| 26.9.xlsx | 61,169行、43,846个归一化客户词、1,515个活动、20个组合 | 本份SP报告的广告花费、归因销售额、投放效率 | 无可靠推广ASIN/SKU字段；每行有覆盖日期区间；含商品编码条目 |
| 对应关系.xlsx | 72个非空且唯一ASIN；其中51个有SQP、21个未导出 | 识别商品、颜色、规格 | 不能证明广告范围仅为这72个ASIN，也不能证明所有都在售 |

“SQP包含自然和广告搜索表现”不等于“这份SQP导出是这份广告报告的总体”。Top100限制的是每个ASIN的查询排名，不是把整个市场按搜索量取前100。词出现在ASIN A的CSV里，也不意味着这个词在ASIN B的记录完整。

原始客户词43,847个，仅将大小写、首尾空格与连续空白统一后变成43,846个；不合并单复数、不同语言和不同拼写。市场计数对相同词完全一致，因此全局只计一次。ASIN计数在已导出ASIN×词范围内求和，称为“样本计数”，不称“完整品牌总量”。

## 2. 数据质量与证据边界

全部跨日期广告行有11,646行、花费$87,876.68，占71.6%。例如开始9月1日、结束9月30日的346行合计花费$41,681.05，这些是覆盖区间计数，不能当作9月1日的发生额。当前不做每日趋势，也不只剔除跨日期行或平均分摊到每天。月度合计用于本份报告的效率分析；完整性仍应与后台同范围月度汇总对账。

{table(evidence,['证据层级','依据','例子','解释边界'])}

**精度问题：** 原始ASIN份额与直接计数相除的差均在0.005个百分点内，符合百分比保留两位小数。不能因为“总量×显示份额”反推时有误差，就断言原始整数计数被脱敏。价格有9行出现约6万至9万的异常值，本次不用于价格分析，也不凭猜测修正。

## 3. SQP：先识别商品承接与相关性

### 3.1 主力商品决定优化优先级

SQP样本共3,416,240曝光、47,159点击、16,730加购、3,778购买计数。整体购买/点击为8.0%。购买计数为ASIN口径，跨ASIN求和不等于独立买家数或去重订单笔数。

{table(asin.head(8),['ASIN','品名','样本点击','样本购买','点击后购买率','结构调整购买指数'],fmt)}

“结构调整购买指数”=实际样本购买÷Σ（每词样本点击×同词市场购买/市场点击）。例如0.77表示按该ASIN已观察到的点击词结构与同词市场率对照，实际购买计数约为基准的77%。它减少了不同ASIN词结构差异的影响，但不控制价格、广告位置、库存、流量人群等差异；不代表颜色导致差异，也不是显著性或因果结论。

64oz黑色值得优先检查，是因为它体量大、影响面广。32oz黑色581点击134购买，购买/点击23.1%；绿色2,434点击121购买，为5.0%。结构调整后两者指数仍为1.77与0.57，值得检查词结构以外的库存、价格、评价、页面、广告位置等因素。不能只凭这两份表认定黑色比绿色更受欢迎。

![商品集中与承接](图1_商品集中与承接.png)

### 3.2 大词不一定值得扩大投入

{table(gaps,['搜索词','样本点击','样本购买','样本点击后购买率','市场点击后购买率'],fmt)}

`kitchen`有2,205样本点击但仅22购买，购买/点击约1.0%；`home`为0.8%。市场本身这类词的购买率也低，因此不能单凭低购买率判定商品页面有问题，也不能因为曝光或点击大就优先投放。原数据没有分出自然点击，不能将这些词定性为“自然大词”。

更值得诊断的是`mason jars`：样本购买/点击3.9%，同词市场12.4%；`cold brew pitcher`为5.8%对12.0%。其中64oz黑色在`mason jars`为1,465点击、57购买，原CSV第62行；在`cold brew pitcher`为960点击、50购买，原CSV第79行。先判断用户想要普通储存罐还是带把手倒水壶、想要冷萃功能是否满足，再查页面承接。没有页面与库存数据，不能确定具体原因。

诊断表里的“距市场率购买差额”用于排序检查顺序；市场竞争商品不同，不能把差额当成可挽回订单或收入承诺。

### 3.3 适配词和商品的增长线索

`glass carafe with lid`市场搜索量8,849，样本145点击、21购买，购买/点击14.5%，同词市场10.6%，样本购买份额4.94%。其中64oz黑色B0DGCVK1JL贡献144点击、21购买（原CSV第63行）。它是一个需商品形态核查的小额测试候选：确认是否符合carafe、带盖等意图后，再测试文案或精准词覆盖。不能将一个月145点击的结果当成稳定增长预测。

样本总体按自身点击词结构加权的市场购买率为9.4%，实际为8.0%，指数0.85。这个对照提示承接值得进一步拆解，不表示整体“少卖了多少订单”。更不能把不同ASIN行的市场总量相加当作全市场规模。

## 4. 广告：先看结构，再看具体词与投放单元

### 4.1 广告整体效率

| 指标 | 9月本份广告报告合计 |
| --- | --- |
| 曝光 / 点击 | 8,039,471 / 158,568 |
| 花费 / 7天归因销售额 | $122,725.96 / $412,829.06 |
| 广告归因订单 / 销量 | 16,016 / 16,571 |
| CTR / CPC / 订单每点击 | 1.97% / $0.77 / 10.10% |
| 总ACOS / ROAS | 29.73% / 3.36 |
| 广告SKU / 其他SKU归因销售额 | $231,125.83 / $181,703.23 |
| 其他SKU销售额占比 | 44.01% |

所有比率均用合计分子÷合计分母，未平均原表逐行比率。零销售额时ACOS留空，不填成0。

ACOS29.7%仅意味着每$100归因销售额对应约$29.7广告花费，是否盈利取决于广告前的贡献利润率：销售收入扣除商品成本、平台佣金、履约、物流、优惠等之后的利润率。退款等也应按业务口径纳入。未提供成本数据，因此不下“盈利/亏损”判断。归因销售额也不等于广告带来的增量销售。

广告花费÷广告SKU销售额为53.1%。这个对照显示对其他SKU归因收入依赖较高，不是亚马逊标准总ACOS的替代指标，不代表其他SKU收益应该忽略；“其他SKU”也不一定是另一条产品线，可能含同款变体。当前表无法确认具体售出商品及是否增量。

### 4.2 广告组合的优先级

{table(port.head(10),['广告组合名称','广告花费','广告销售额','广告订单','ACOS','其他SKU销售额占比'],fmt)}

【梅森罐】占花费66.6%、销售额69.7%，总ACOS28.4%，是最值得持续维护的主力组合。【32oz梅森罐】ACOS33.0%，其他SKU收入占57.0%；【果茶款梅森罐】ACOS37.3%，其他SKU占56.4%；应同时拆看成本和实际售出商品。组合名称只是后台分组，不自动等于某个ASIN或完整产品线。

### 4.3 同一词要看投给了什么商品

{table(same,['广告组合名称','广告花费','广告订单','ACOS'],fmt)}

`mason jar pitcher`在不同组合都出单，但成本明显不同。【梅森罐】22.9%可列增长测试，【32oz梅森罐】47.3%先列降本复核。建议先核对具体投放单元和推广商品，再考虑组合内小幅调整；不能由组合名直接判断原因是容量。

![同词不同组合](图3_同词不同组合.png)

### 4.4 不要把所有“-”算自动广告

{table(types[types['广告花费']>=800],['投放类型','广告花费','广告订单','ACOS','其他SKU销售额占比'],fmt)}

四类自动通过投放字段识别；“-”还包含商品ASIN定位、类目与关键词组定向。自动关联ACOS46.3%、其他SKU收入占62.7%，自动同类36.7%，优先查看具体单元。手动词组26.6%低于精准29.4%，但商品和词构成不同，不能得出“词组优于精准”的普遍规则。

具体有体量的问题单元：

- 【梅森罐】低价扫流0.7 -2 - z，`substitutes`：$3,654.70、346单、ACOS41.9%。先拆客户ASIN与推广商品，再决定哪部分调整。
- 【32oz梅森罐】低价扫流0.7 - z，`complements`：$944.69、77单、ACOS53.3%。先核对关联商品和连带利润，再小步测试成本。

![投放类型效率](图4_投放类型效率.png)

### 4.5 零单花费不能全部叫浪费

原始行中零单花费$50,472.72；同词在其他行可能有成交。按“广告组合×客户词”汇总后，整个组合词仍零单的花费为$45,025.65，大量来自一两次点击的条目。按“≥20点击、≥$20且零单”筛选仅26个组合词、$702.04，占全部花费0.57%。因此直接声称可以节省$50,472.72没有依据。

{table(zero,['广告组合名称','搜索词','广告点击','广告花费'],fmt)}

20点击零单只是复核阈值：如果真实购买率10%，20点击仍零单的概率约12%。核对归因是否成熟、相关性和具体投放单元后才考虑否定；文本词与商品编码分别处理，不能把整个店铺都否掉。

### 4.6 有单高成本与增长候选

复核筛选：≥5单、花费≥$50且ACOS>40%，共45个组合词、$6,945.61花费，均有成交，适合检查成本和利润，不适合直接暂停。

{table(weak,['广告组合名称','搜索词','广告花费','广告订单','ACOS'],fmt)}

增长筛选：≥10单且ACOS≤25%，共105个组合词，合计花费$18,854.22、归因销售额$91,354.01、3,625单。25%和40%是本次用于筛选的观察线，用户暂未提供业务目标，不能当作统一盈亏标准。

{table(grow,['广告组合名称','搜索词','广告花费','广告订单','ACOS'],fmt)}

具体可讨论的投放单元：【梅森罐】手动精准核心1 mason jar pitcher，花费$2,131.72、374单、ACOS22.3%；以及【梅森罐】手动广泛 glass pitchers with handle and lid - z，$4,029.91、633单、23.6%。先看推广商品、预算是否受限、库存和目标ACOS；预算受限才试加预算，避免因为历史效率好就机械加价。更高预算的边际成本可能上升。

## 5. 两表对照：为什么“全量小于纯广告”

| 集合 | 词条数 | SQP样本点击 | 广告点击 | 广告点击占全表 |
| --- | --- | --- | --- | --- |
| 两表都收录 | 1,802 | 46,005 | 56,382 | 35.56% |
| 仅SQP收录 | 608 | 1,154 | 未匹配到 | — |
| 仅广告收录 | 42,044 | 未匹配到 | 102,186 | 64.44% |

1,802+608=2,410个SQP词；1,802+42,044=43,846个广告词；并集44,454个。以双方是否出现做集合统计，与双方是否有正点击分开。1,776个双方有点击只是交集子集，另外26个共有词在SQP样本中零点击。

仅广告收录的42,044条中，10,681条是ASIN编码，贡献46,191点击；31,363条是文本词，贡献55,995点击。编码条目不宜解释为普通客户输入的“长尾词”，是否发生在详情页或其他位置仍需广告报告确认。文本词也可能来自未导出商品，不能把全部缺失都归因于Top100。

![广告点击覆盖](图2_广告点击覆盖.png)

即使只看共有词，广告点击56,382仍大于SQP样本46,005，说明按词取交集仍未解决ASIN与场景范围差异。广告销售包含其他SKU、购买归因时间规则也需对齐。因此不计算自然点击、自然购买或广告占比；广告≤SQP时也不使用相减。仅SQP出现也只表示本份SP报告未匹配到，不意味着完全没广告。

## 6. 给带教的执行建议与下一步数据

### 第一步：先校准口径和主力商品

本次以可复算的月度总数和原始SQP记录为基础；日度趋势与自然流量拆分需补齐口径后再评估。核对64oz黑色的`mason jars`与`cold brew pitcher`搜索意图、页面、库存和实际投放商品。补充广告前贡献利润率及不同SKU毛利，建立各商品合理目标ACOS。

### 第二步：降本和增长分别做小步测试

降本从上述同类/关联单元、45个有单高ACOS组合词里选择少量具体对象；增长从105个候选中选择库存充足、利润符合要求且受预算约束的对象。作为讨论方案，单次出价可试降10%—15%、预算可试增10%，每次只调整一个变量。这些幅度不是对全部单元的执行指令。

记录修改日期、原设置、商品、目标、预期观察指标。先观察7天点击发生数据，再等待7天归因窗口及后台更新充分成熟后比较；需覆盖可比星期结构，避免同时更改价格、图片或优惠。观察ACOS、订单/日、广告SKU与其他SKU收入、全店或商品总销量，不只看单一广告比例。低点击测试延长观察，不机械按7天结束。

### 第三步：把补数据请求写具体

- SQP：同市场、同月份、同ASIN范围，申请后台当前权限允许导出的更多词；如能取品牌视图，保留品牌字段及完整指标说明。不要预设一定能导出1,000条。
- 广告：重新导出明确日粒度SP报告；另取推广商品报告，将活动/广告组与推广ASIN、SKU、售出商品建立可靠关系。若广告组含多个商品，不能只凭广告组名称硬分摊搜索词指标。
- 商品经营：广告前贡献利润率、库存与断货记录、价格和优惠、评论变化、Business Reports商品销售与Sessions、其他月份同口径数据。
- 归因说明：保留当前后台SQP购买定义、广告7天窗口、更新日期和报告设置截图，再决定可比较范围。补充SB/SD是否存在，不能默认本份SP就是所有广告。

## 7. 文件和复算方式

优先阅读《给带教的汇报稿.md》，核对时打开《9月独立分析.xlsx》。完整CSV包括全部51,736条组合词汇总、43,846条广告客户词汇总和44,454条两表对照，没有为了美观删除低点击词。工作簿只放与SQP相关的2,410条对照和176条行动候选，其余见同目录CSV。

SQP明细保留来源文件与原始行号；广告组合词汇总保留原始Excel行号和日期范围。原始行号首末只用于定位范围，不代表中间行均匹配。原始文件SHA-256记录于source_manifest.json。

复算：在项目根目录运行 `python analysis/analyze.py`，再运行 `python analysis/build_deliverables.py`。只读取原始CSV/XLSX，不读取其他分析文件作为数据来源，不改变原始压缩包或raw_data。
'''
matrix_file=OUT/'9月_搜索机会矩阵.csv'
if matrix_file.exists():
    report+='\n\n## 8. 搜索机会矩阵与使用场景\n\n点击吸引力指数＝样本点击/样本曝光 ÷ 同词市场点击/市场曝光；购买承接指数＝样本购买/样本点击 ÷ 同词市场购买/市场点击。以1倍作为描述性分界，低于100点击先标为样本不足。这个矩阵用于检查顺序，不是显著性检验，不识别自然流量，也不证明页面因素导致差异。\n\n| 点击吸引力 | 购买承接 | 优先检查 |\n| --- | --- | --- |\n| 达基准 | 达基准 | 相关性、绝对购买率、利润、库存和预算约束，之后考虑维护或增长测试 |\n| 达基准 | 偏低 | 规格和词意图、价格、页面、库存、评价、流量位置 |\n| 偏低 | 达基准 | 主图、展示位置、标题、价格和竞品构成 |\n| 偏低 | 偏低 | 商品与词意图是否适配，避免只因市场词大就买量 |\n\n场景线索：sun tea glass jar for outside的吸引力指数1.47、承接指数1.14，【梅森罐】广告ACOS20.3%；tea pitcher for fridge分别为1.48、1.15与19.6%。可以结合商品实际支持的容量、盖子、把手、倒水方式考虑小额增长测试。\n\niced tea pitcher的SQP两项指数也较强，但【梅森罐】广告ACOS31.2%，先看利润与成本。glass carafe with lid承接指数1.37、吸引力0.88、广告ACOS13.3%，可先研究点击环节；只有145个SQP样本点击，仍需验证稳定性。上述两表范围未完全对齐，不将广告表现归到某个具体ASIN，也不推算自然流量。\n\n7月工作文件仅提供分析组织方式参考，其数值、第三方排名、竞价与公式没有作为9月指标来源。矩阵和场景判断均用9月原始数据重新计算。\n'
(OUT/'9月独立分析报告.md').write_text(report)

summary='''# 给带教的汇报稿

我从9月原始SQP、广告搜索词报告和商品对应关系重新计算，主要分析商品搜索承接、广告降本与增长机会。结论分为可验证事实、诊断线索和待验证建议。

首先，两份数据目前不能直接相减。SQP只有51个ASIN导出的前100查询样本，广告报告覆盖的商品和场景也不完全一致。SQP没出现的广告条目有42,044个，其中10,681个是ASIN编码，不能全部叫长尾搜索词。广告原表还有很多跨日期区间，按开始日期做每日趋势会失真，所以这次先用月度汇总分析。

我主要从降本和增长两个方向看。

**商品这边，建议先看64oz黑色单装梅森罐。** 它占SQP样本70.5%的点击、61.9%的购买计数，影响最大。它在mason jars的购买/点击是3.9%，同词市场为12.4%；cold brew pitcher是5.2%对12.0%。我想先检查搜索意图是否适配商品，以及主图、功能说明、价格、评价和库存，暂时不能确定是哪一项导致差异。

**广告这边，同一个词要拆到商品组合和投放单元。** mason jar pitcher在【梅森罐】组合ACOS22.9%，在【32oz梅森罐】47.3%。前者可以列增长候选，后者先查成本和落地商品，不能全店对这个词统一加价或否词。

**降本优先看有体量的具体单元。** 【梅森罐】低价扫流0.7 -2 - z的同类投放，花费$3,654.70、346单、ACOS41.9%；【32oz梅森罐】低价扫流0.7 - z的关联投放，$944.69、77单、53.3%。建议先拆目标再小步测试。零单词按“20点击且$20以上”筛出26个，合计仅$702.04，不适合直接把原始零单行花费$50,472.72都算成浪费。

**增长可先维护有验证体量的词。** 【梅森罐】手动精准核心1 mason jar pitcher，374单、ACOS22.3%，可以在确认利润、库存和预算限制后，考虑小幅增加预算；先不要同时加出价，避免不知道变化来自哪里。

另一个需要一起看的是销售构成：9月广告花费$122,725.96、归因销售额$412,829.06，总ACOS29.7%，但44.0%的归因销售额来自其他SKU。这些收入有价值，不过目前不能确认具体售出商品，也不能用总ACOS直接判断被投商品盈利。

下一步我建议先补目标ACOS或广告前贡献利润率、可靠的推广ASIN映射和日粒度广告报告。然后挑少量高成本单元做降本测试，同时给少量高效词做增长测试，等归因成熟后看成本、订单和商品总销量。详细证据、候选和原始行号已放在独立分析工作簿里。
'''
if matrix_file.exists():
    summary+='\n我补了同词“点击吸引力×购买承接”的机会矩阵。它提示太阳茶、冰箱茶壶等明确场景可以继续验证；但搜索表现好仍要结合广告成本，比如iced tea pitcher搜索表现较强，广告ACOS却为31.2%，不能只看一个维度就加预算。7月文件作为思路参考，没有并入9月数据。\n'
(OUT/'给带教的汇报稿.md').write_text(summary)

# Workbook: values are precomputed so opening without Excel recalculation is reliable.
sqp_cross=cross[cross['集合状态'].isin(['两表都收录','仅SQP收录'])].sort_values('样本点击',ascending=False)
use_cross=['搜索词','集合状态','收录ASIN数','市场搜索量','样本曝光','样本点击','样本加购','样本购买',
    '广告曝光','广告点击','广告花费','广告销售额','广告订单','ACOS','解释边界']
workbook_tables={
 '00_说明':notes,'01_证据层级':evidence,'02_核心行动':actions,
 '03_广告总览':pd.DataFrame([M['ad_totals']]),'04_广告组合':port,'05_投放类型':types,'06_客户词类型':qtypes,
 '07_SQP商品':asin,'08_SQP词诊断':qs,'09_两表覆盖':coverage,
 '10_两表对照_SQP相关':sqp_cross[use_cross], '11_零单待复核':queues[queues['建议队列'].str.startswith('优先')],
 '12_高ACOS待复核':queues[queues['建议队列'].str.startswith('复核')],
 '13_增长候选':queues[queues['建议队列'].str.startswith('增长')],
 '14_投放单元':targets,'15_SQP来源明细':aq,'16_质量检查':checks,
 '17_价格异常':read('质量_价格异常'),'18_未导出ASIN':read('范围_未导出ASIN'),
}
if matrix_file.exists():
    workbook_tables['19_搜索机会矩阵']=read('9月_搜索机会矩阵')
with pd.ExcelWriter(OUT/'9月独立分析.xlsx',engine='openpyxl') as writer:
 for name,d in workbook_tables.items(): d.to_excel(writer,sheet_name=name,index=False)
 for ws in writer.book:
    nrows=ws.max_row
    small_sheet=nrows<250
    ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions;ws.sheet_view.showGridLines=False
    ws.row_dimensions[1].height=34
    for cell in ws[1]:
        cell.fill=PatternFill('solid',fgColor='173D59');cell.font=Font(name='微软雅黑',color='FFFFFF',bold=True,size=10)
        cell.alignment=Alignment(vertical='center',wrap_text=True)
    headers={c.column:c.value for c in ws[1]}
    for j,h in headers.items():
        if h in ['解释','问题或核对点','复核证据','建议处理','建议测试或核查','观察指标及边界','解释边界','参考文件证据或内容','本次如何使用','建议关注','判断边界']: width=60
        elif h in ['广告活动名称','广告组名称','投放','原始搜索词','搜索词','品名','任务','对象','证据','已有文件位置']: width=38
        else: width=max(13,min(28,len(str(h))*2+3))
        ws.column_dimensions[get_column_letter(j)].width=width
    for row in ws.iter_rows(min_row=2):
        for c in row:
            h=headers[c.column];c.font=Font(name='微软雅黑',size=10)
            if isinstance(c.value,str):
                c.data_type='s' # CSV-derived strings cannot become executable Excel formulas.
                c.alignment=Alignment(vertical='top',wrap_text=small_sheet)
            elif isinstance(c.value,(int,float)):
                if h in ['CTR','CVR','ACOS','广告SKU销售额口径ACOS','其他SKU销售额占比','花费占比','销售额占比','样本购买贡献占比','广告点击占全表比','点击份额','购买份额'] or ('率' in str(h) and '预期' not in str(h)):
                    c.number_format='0.00%'
                elif any(x in str(h) for x in ['花费','销售额','CPC']):c.number_format='"$"#,##0.00'
                elif any(x in str(h) for x in ['指数','ROAS','每加购比','差额','预期购买']):c.number_format='0.00'
                else:c.number_format='#,##0'
            if isinstance(c.value,pd.Timestamp):c.number_format='yyyy-mm-dd'
        if small_sheet:
            wrapped_lines=1
            for c in row:
                if isinstance(c.value,str):
                    display_width=sum(2 if ord(ch)>255 else 1 for ch in c.value)
                    col_width=ws.column_dimensions[get_column_letter(c.column)].width
                    wrapped_lines=max(wrapped_lines,math.ceil(display_width/max(col_width-2,1)))
            ws.row_dimensions[row[0].row].height=max(32,min(170,wrapped_lines*15+8))
    for j,h in headers.items():
        if h=='ACOS' and nrows>2:
            col=get_column_letter(j)
            ws.conditional_formatting.add(f'{col}2:{col}{nrows}',ColorScaleRule(start_type='num',start_value=.2,start_color='DCEFE8',end_type='num',end_value=.6,end_color='F8D2C6'))
 # Add a native Excel chart for reviewers to inspect its source values.
 ws=writer.book['04_广告组合']; ch=BarChart();ch.title='广告组合总ACOS（前10个花费组合）';ch.y_axis.title='ACOS'
 ac=list(port.columns).index('ACOS')+1
 ch.add_data(Reference(ws,min_col=ac,min_row=1,max_row=11),titles_from_data=True)
 ch.set_categories(Reference(ws,min_col=1,min_row=2,max_row=11));ch.height=10;ch.width=22
 ws.add_chart(ch,'V2')

# Word reports include the same tables and figures; no remote publishing.
def make_doc(md,path):
 doc=Document();sec=doc.sections[0];sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.65)
 sec.left_margin=Inches(.7);sec.right_margin=Inches(.7)
 st=doc.styles['Normal'];st.font.name='Calibri';st.font.size=Pt(10)
 st._element.rPr.rFonts.set(qn('w:eastAsia'),'微软雅黑')
 st.paragraph_format.space_after=Pt(6)
 for name in ['Title','Heading 1','Heading 2','Heading 3']:
    doc.styles[name]._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'微软雅黑')
 def clean(s):return s.replace('**','').replace('`','')
 lines=md.splitlines();i=0
 while i<len(lines):
    line=lines[i].strip()
    if not line:i+=1;continue
    if line.startswith('|'):
        block=[]
        while i<len(lines) and lines[i].strip().startswith('|'):
            block.append([clean(x.strip()) for x in lines[i].strip().strip('|').split('|')]);i+=1
        rows=[block[0]]+block[2:]; t=doc.add_table(rows=0,cols=len(rows[0]));t.style='Light Shading Accent 1'
        for rr in rows:
            cells=t.add_row().cells
            for c,v in zip(cells,rr):
                c.text=v
                for p in c.paragraphs:
                    for run in p.runs:run.font.size=Pt(8)
        continue
    match=re.match(r'!\[(.*?)\]\((.*?)\)',line)
    if match:
        doc.add_picture(str(OUT/match.group(2)),width=Inches(6.7));doc.add_paragraph(match.group(1));i+=1;continue
    if line.startswith('# '): doc.add_heading(clean(line[2:]),0)
    elif line.startswith('## '): doc.add_heading(clean(line[3:]),1)
    elif line.startswith('### '): doc.add_heading(clean(line[4:]),2)
    elif line.startswith('- '): doc.add_paragraph(clean(line[2:]),style='List Bullet')
    elif re.match(r'^\d+\. ',line):doc.add_paragraph(clean(re.sub(r'^\d+\. ','',line)),style='List Number')
    else: doc.add_paragraph(clean(line))
    i+=1
 footer=sec.footer.paragraphs[0];footer.alignment=2
 footer.add_run('2026年9月｜独立复核    ')
 field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
 doc.save(path)
make_doc(report,OUT/'9月独立分析报告.docx')
make_doc(summary,OUT/'给带教的汇报稿.docx')

readme='''# 交付文件说明

先看：给带教的汇报稿.md / .docx（可用于口头汇报）。

完整报告：9月独立分析报告.md / .docx。

核对数据：9月独立分析.xlsx，20个页签，包含证据层级和搜索机会矩阵。所有结果写为数值，无需等待公式重算。

四张PNG图可独立插入汇报。完整细节保留在CSV中：广告_组合词汇总（51,736行）、广告_全部词汇总（43,846行）、两表_全部词对照（44,454行）等。

复算需Python与pandas、numpy、openpyxl、matplotlib、python-docx；本环境已安装。项目根目录依次执行：

```bash
python analysis/analyze.py
python analysis/build_search_matrix.py
python analysis/build_deliverables.py
```

原始文件位于source_material/raw_data，是原zip解包副本。source_manifest.json记录原始文件哈希。所有指标仅由raw_data中的CSV与Excel原表计算。原始文件未改动。

7月工作文件仅作为组织思路参考，未作为9月指标来源。搜索机会矩阵由build_search_matrix.py直接使用9月独立计算结果生成。

重要口径：缺失不等于0；SQP样本购买计数不是全店订单；不求自然流量差值；不把日期区间当日数据；广告组合不等于ASIN；目标ACOS尚未设定，行动阈值仅筛选待复核候选。
'''
(OUT/'README.md').write_text(readme)
print('Independent report, Word, workbook and figures created in',OUT)
