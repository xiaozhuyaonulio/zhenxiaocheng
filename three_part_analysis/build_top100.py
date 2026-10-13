"""Standalone SQP click-top-100 result; no advertising data or project bundle."""
from pathlib import Path
import hashlib,json,re,zipfile,math
from lxml import etree as ET
import numpy as np
import pandas as pd
import openpyxl
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties

ROOT=Path(__file__).resolve().parents[1]
INPUT=ROOT/'reference/26.9汇总表_三块分析底稿.xlsx'
OUT=ROOT/'three_part_analysis'
BOOK=OUT/'SQP_点击前100词分析.xlsx'
norm=lambda x:re.sub(r'\s+',' ',str(x).strip().lower())
source=pd.read_excel(INPUT,sheet_name='原始SQP合并')
source['搜索词']=source['搜索查询'].map(norm)
source['原稿来源行']=np.arange(2,len(source)+2)
own={'曝光量：ASIN 计数':'样本曝光','点击量：ASIN 数量':'样本点击',
     '加购：ASIN 数量':'样本加购','下单成交：ASIN 数量':'样本购买'}
market={'搜索查询量':'市场搜索量','曝光：曝光总量':'市场曝光','点击量：总次数':'市场点击',
        '购物车添加：总数':'市场加购','购买：下单总数':'市场购买'}
assert not source.duplicated(['ASIN','搜索词']).any()
assert not source.groupby('搜索词')[list(market)].nunique(dropna=False).gt(1).any().any()
q=source.groupby('搜索词')[list(own)].sum().rename(columns=own)
q=q.join(source.drop_duplicates('搜索词').set_index('搜索词')[list(market)].rename(columns=market))
q=q.join(source.groupby('搜索词').agg(收录ASIN数=('ASIN','nunique'),原稿SQP首行=('原稿来源行','min'))).reset_index()
pairs={'样本CTR':('样本点击','样本曝光'),'市场CTR':('市场点击','市场曝光'),
       '样本加购点击比':('样本加购','样本点击'),'样本购买点击比':('样本购买','样本点击'),
       '市场购买点击比':('市场购买','市场点击')}
for name,(n,d) in pairs.items():q[name]=q[n]/q[d].replace(0,np.nan)
q['CTR差_百分点']=(q['样本CTR']-q['市场CTR'])*100
q['购买比差_百分点']=(q['样本购买点击比']-q['市场购买点击比'])*100
q['市场率参考购买']=q['样本点击']*q['市场购买点击比']
tot=q[['样本曝光','样本点击','样本加购','样本购买']].sum()
q['点击贡献占比']=q['样本点击']/tot['样本点击']
q['购买贡献占比']=q['样本购买']/tot['样本购买']
primary=source.sort_values(['搜索词','点击量：ASIN 数量','ASIN'],ascending=[True,False,True]).drop_duplicates('搜索词')
q=q.merge(primary[['搜索词','ASIN','点击量：ASIN 数量']].rename(
    columns={'ASIN':'主要点击ASIN','点击量：ASIN 数量':'主要ASIN点击'}),on='搜索词',validate='one_to_one')
q['主要ASIN点击占比']=q['主要ASIN点击']/q['样本点击'].replace(0,np.nan)

def describe(r):
    if r['样本点击']<100:return '点击不足100，暂不作购买表现判断','先检查曝光与词意图；继续观察购买计数'
    if pd.isna(r['市场率参考购买']) or r['市场率参考购买']<10:
        return '购买基准计数少，先观察','核查词意图和主力ASIN；不凭比率判断购买优劣'
    a,b=r['CTR差_百分点']<0,r['购买比差_百分点']<0
    if a and b:return 'CTR及购买比均低于市场（描述）','先查词意图与主力ASIN，再核查商品和展示条件'
    if b:return 'CTR不低，购买比低于市场（描述）','查主力ASIN规格、功能、价格、库存及页面；原因待验证'
    if a:return 'CTR低，购买比不低于市场（描述）','核查展示位置、主图标题、价位及竞争构成'
    return 'CTR及购买比均不低于市场（描述）','确认相关性和表现稳定性；增长测试还需成本、利润与库存'
descriptions=q.apply(describe,axis=1)
q['分析线索']=[x[0] for x in descriptions]
q['下一步检查']=[x[1] for x in descriptions]
top=q.sort_values(['样本点击','搜索词'],ascending=[False,True]).head(100).copy()
expo=q.sort_values(['样本曝光','搜索词'],ascending=[False,True]).head(100).copy()
extra=expo[~expo['搜索词'].isin(top['搜索词'])].copy()
top.insert(0,'点击排名',range(1,101))
expo_rank={word:i for i,word in enumerate(expo['搜索词'],1)}
top['曝光排名_仅曝光前100']=top['搜索词'].map(expo_rank)
extra.insert(0,'曝光排名',[expo_rank[x] for x in extra['搜索词']])
display=['搜索词','市场搜索量','收录ASIN数','样本曝光','样本点击','点击贡献占比','样本加购','样本购买',
         '购买贡献占比','样本CTR','市场CTR','CTR差_百分点','样本加购点击比','样本购买点击比',
         '市场购买点击比','购买比差_百分点','主要点击ASIN','主要ASIN点击','主要ASIN点击占比',
         '市场曝光','市场点击','市场购买','市场率参考购买','分析线索','下一步检查','原稿SQP首行']
main=top[['点击排名']+display+['曝光排名_仅曝光前100']]
extra=extra[['曝光排名']+display]
cs=top[['样本曝光','样本点击','样本加购','样本购买']].sum()
es=expo[['样本曝光','样本点击','样本购买']].sum()
pct=lambda x:f'{x:.2%}'
summary=pd.DataFrame([
 ['范围','仅本次SQP源页；51个ASIN、3,117条ASIN×词记录、2,410个去重词。','不是商品完整流量；不使用广告数据。'],
 ['点击前100覆盖',f'{int(cs["样本点击"]):,}/{int(tot["样本点击"]):,}点击={pct(cs["样本点击"]/tot["样本点击"])}；{int(cs["样本购买"]):,}/{int(tot["样本购买"]):,}购买={pct(cs["样本购买"]/tot["样本购买"])}。','同一组点击前100词；不是分别选择购买前100。'],
 ['曝光前100覆盖',f'{int(es["样本曝光"]):,}/{int(tot["样本曝光"]):,}曝光={pct(es["样本曝光"]/tot["样本曝光"])}；点击覆盖{pct(es["样本点击"]/tot["样本点击"])}；购买覆盖{pct(es["样本购买"]/tot["样本购买"])}。','曝光排序适合查展示规模及曝光后点击表现。'],
 ['两种榜单的差异',f'91词同时进入两榜；9词只进入曝光榜，见曝光补充词页。','点击榜适合研究已有点击后的行为；单看点击榜会漏掉部分高曝光低点击词。'],
 ['选词规则','样本点击降序；并列时搜索词升序。第100名为coffee container，101点击、4购买；glass carafe也为101点击，按并列规则排第101。','原始每ASIN前100与这张合并点击前100是不同概念；边界并列不表示表现优劣。'],
 ['计数支持','点击前100各词至少101点击；购买表现筛查另看市场率参考购买是否达到10。','100点击和10参考购买均为工作筛选线；不代表显著性或稳定性。'],
 ['结论使用','CTR和购买比差只描述同词样本与市场差异；分析线索提供核查方向。','不按倍率排序，不将差额当可挽回订单；不能单凭SQP确定利润、广告效果或自然量。'],
],columns=['问题','结果或说明','解释边界'])
lookup=q.set_index('搜索词')
for word in ['mason jars','cold brew pitcher','tea pitcher for fridge']:
    r=lookup.loc[word]
    summary.loc[len(summary)]=[word,
       f'{int(r["样本点击"]):,}点击/{int(r["样本购买"]):,}购买；CTR样本{pct(r["样本CTR"])}、市场{pct(r["市场CTR"])}；购买/点击样本{pct(r["样本购买点击比"])}、市场{pct(r["市场购买点击比"])}。',
       r['分析线索']+'；'+r['下一步检查']]

def addr(cols,h,row):return get_column_letter(cols.index(h)+1)+str(row)
def formula(h,cols,row):
    a=lambda x:addr(cols,x,row)
    if h in pairs:
        n,d=pairs[h]
        return f'=IFERROR(IF(AND(ISNUMBER({a(n)}),ISNUMBER({a(d)}),{a(d)}>0),{a(n)}/{a(d)},""),"")'
    if h in ('点击贡献占比','购买贡献占比'):
        n,total='样本点击' if h=='点击贡献占比' else '样本购买',int(tot['样本点击' if h=='点击贡献占比' else '样本购买'])
        return f'={a(n)}/{total}'
    if h=='主要ASIN点击占比':return f'=IFERROR({a("主要ASIN点击")}/{a("样本点击")},"")'
    if h in ('CTR差_百分点','购买比差_百分点'):
        n,d=('样本CTR','市场CTR') if h=='CTR差_百分点' else ('样本购买点击比','市场购买点击比')
        return f'=IFERROR(IF(AND(ISNUMBER({a(n)}),ISNUMBER({a(d)})),({a(n)}-{a(d)})*100,""),"")'
    if h=='市场率参考购买':return f'=IFERROR(IF(ISNUMBER({a("市场购买点击比")}),{a("样本点击")}*{a("市场购买点击比")},""),"")'
    return None
guide=[]
for h in main.columns:
    f=formula(h,list(main.columns),2)
    guide.append([h,f or '原始取数、按词汇总、排序或公开规则标签；不另造数值公式。',
                  '贡献占比分母固定为全部2,410词；源版本变更需重算。' if '贡献占比' in h else '见阅读说明；原稿SQP首行对应底稿源页。'])
guide=pd.DataFrame(guide,columns=['字段','主表第2行公式或取数','说明'])
frames={'阅读说明':summary,'点击前100分析':main,'曝光补充9词':extra,'字段公式':guide}
wb=openpyxl.Workbook(write_only=True);wb.calculation=CalcProperties(calcMode='auto',fullCalcOnLoad=False)
caches={};sizes={}
for name,df in frames.items():
    ws=wb.create_sheet(name);ws.freeze_panes='C2' if len(df.columns)>10 else 'A2';ws.sheet_view.showGridLines=False
    cols=list(df.columns);ws.row_dimensions[1].height=48
    for j,h in enumerate(cols,1):
        ws.column_dimensions[get_column_letter(j)].width=65 if h in ('结果或说明','解释边界','主表第2行公式或取数','说明','下一步检查') else 40 if h in ('搜索词','分析线索') else 22
    head=[]
    for h in cols:
        c=WriteOnlyCell(ws,h);c.font=Font(name='微软雅黑',bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='173D59');c.alignment=Alignment(wrap_text=True,vertical='center');head.append(c)
    ws.append(head);caches[name]={}
    for row,values in enumerate(df.itertuples(index=False,name=None),2):
        if len(df)<30:ws.row_dimensions[row].height=65
        cells=[]
        for j,(h,v) in enumerate(zip(cols,values),1):
            v=None if pd.isna(v) else v.item() if isinstance(v,np.generic) else v
            f=formula(h,cols,row) if name in ('点击前100分析','曝光补充9词') else None
            c=WriteOnlyCell(ws,f or v)
            if f:caches[name][f'{get_column_letter(j)}{row}']=v
            elif isinstance(v,str):c.data_type='s'
            c.font=Font(name='微软雅黑',size=10);c.alignment=Alignment(vertical='top',wrap_text=len(df)<30 or h in ('分析线索','下一步检查'))
            c.number_format='0.00%' if h in pairs or h.endswith('占比') else '0.00' if '百分点' in h or h=='市场率参考购买' else '#,##0' if isinstance(v,(int,float)) else 'General'
            cells.append(c)
        ws.append(cells)
    sizes[name]=(len(df)+1,len(cols));ws.auto_filter.ref=f'A1:{get_column_letter(len(cols))}{len(df)+1}'
temp=BOOK.with_suffix('.uncached.xlsx');wb.save(temp)
NS='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
with zipfile.ZipFile(temp) as src,zipfile.ZipFile(BOOK,'w',zipfile.ZIP_DEFLATED,compresslevel=6) as dst:
    for i in src.infolist():
        data=src.read(i.filename)
        m=re.fullmatch(r'xl/worksheets/sheet(\d+)\.xml',i.filename)
        if m:
            name=list(frames)[int(m[1])-1];xml=ET.fromstring(data)
            for c in xml.findall(f'.//{{{NS}}}sheetData/{{{NS}}}row/{{{NS}}}c'):
                if c.attrib['r'] in caches[name]:
                    value=caches[name][c.attrib['r']];v=c.find(f'{{{NS}}}v')
                    if v is None:v=ET.SubElement(c,f'{{{NS}}}v')
                    v.text=str(value) if value is not None else ''
                    if value is None:c.attrib['t']='str'
            nr,nc=sizes[name];xml.insert(1,ET.Element(f'{{{NS}}}dimension',{'ref':f'A1:{get_column_letter(nc)}{nr}'}))
            data=ET.tostring(xml,encoding='utf-8')
        dst.writestr(i,data)
temp.unlink()
assert len(top)==100 and cs['样本点击']==38093 and cs['样本购买']==2645
assert len(extra)==9 and len(set(top['搜索词'])&set(expo['搜索词']))==91
check=openpyxl.load_workbook(BOOK,read_only=True,data_only=True)
for name,df in [('点击前100分析',main),('曝光补充9词',extra)]:
    rows=check[name].iter_rows(values_only=True);heads=next(rows)
    for expected,values in zip(df.itertuples(index=False,name=None),rows):
        for h,a,b in zip(heads,values,expected):
            if pd.isna(b):assert a in (None,'')
            elif isinstance(b,(float,int,np.number)):assert math.isclose(a,float(b),rel_tol=1e-11,abs_tol=1e-11),(h,a,b)
            else:assert a==b
    assert check[name].max_row==len(df)+1
check.close()
f=openpyxl.load_workbook(BOOK,read_only=True,data_only=False)
assert f['点击前100分析'].cell(2,list(main.columns).index('样本CTR')+1).data_type=='f'
f.close()
print(json.dumps({'file':str(BOOK),'bytes':BOOK.stat().st_size,'click_top100_clicks':int(cs['样本点击']),
 'click_top100_purchases':int(cs['样本购买']),'exposure_only_words':len(extra),
 'source_sha256':hashlib.sha256(INPUT.read_bytes()).hexdigest(),'verification':'passed'},ensure_ascii=False))
