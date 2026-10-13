"""Verify scope separation, raw-data preservation, aggregates and workbook formulas."""
from pathlib import Path
import hashlib,json,re,zipfile,posixpath,math
import numpy as np
import pandas as pd
import openpyxl
from lxml import etree as ET

ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'three_part_analysis'
INPUT=ROOT/'reference/26.9汇总表_三块分析底稿.xlsx'
BOOK=OUT/'26.9汇总表_三块分析版.xlsx'
NS='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
REL='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
MC='http://schemas.openxmlformats.org/markup-compatibility/2006'
norm=lambda x:re.sub(r'\s+',' ',str(x).strip().lower())
def resolve(x):return x.lstrip('/') if x.startswith('/') else posixpath.normpath('xl/'+x)
def sheets(z):
    rels={x.attrib['Id']:resolve(x.attrib['Target']) for x in ET.fromstring(z.read('xl/_rels/workbook.xml.rels'))}
    return {x.attrib['name']:rels[x.attrib[f'{{{REL}}}id']] for x in ET.fromstring(z.read('xl/workbook.xml')).find(f'{{{NS}}}sheets')}

metrics=json.loads((OUT/'analysis_metrics.json').read_text())
assert hashlib.sha256(INPUT.read_bytes()).hexdigest()==metrics['source_sha256']
with zipfile.ZipFile(INPUT) as src,zipfile.ZipFile(BOOK) as dst:
    assert dst.testzip() is None
    orig=sheets(src);new=sheets(dst)
    assert len(new)==32 and list(new)[0]=='分析入口'
    root=ET.fromstring(dst.read('xl/workbook.xml'))
    names=list(new)
    for item in root.findall(f'.//{{{NS}}}definedName'):
        if 'localSheetId' in item.attrib:
            scoped=names[int(item.attrib['localSheetId'])]
            assert item.text.startswith(scoped+'!') or item.text.startswith("'"+scoped+"'!"),(scoped,item.text)
    for name,path in orig.items():
        assert src.read(path)==dst.read(new[name]),f'Original worksheet changed: {name}'
    # All original style records keep their indices; shared strings are untouched.
    assert src.read('xl/sharedStrings.xml')==dst.read('xl/sharedStrings.xml')
    oldstyle=ET.fromstring(src.read('xl/styles.xml'));style=ET.fromstring(dst.read('xl/styles.xml'))
    for tag in ('fonts','fills','borders','cellStyleXfs','cellXfs'):
        before=oldstyle.find(f'{{{NS}}}{tag}');after=style.find(f'{{{NS}}}{tag}')
        for a,b in zip(before,after):
            assert ET.tostring(a,method='c14n',exclusive=True)==ET.tostring(b,method='c14n',exclusive=True),(tag,'original style changed')
    for path in ('xl/workbook.xml','xl/styles.xml'):
        tree=ET.fromstring(dst.read(path))
        for node in tree.iter():
            for prefix in node.attrib.get(f'{{{MC}}}Ignorable','').split():
                assert prefix in node.nsmap,(path,'missing ignorable namespace',prefix)
            if node.tag==f'{{{MC}}}Choice':
                for prefix in node.attrib.get('Requires','').split():assert prefix in node.nsmap
    style_count=len(style.find(f'{{{NS}}}cellXfs'))
    for name,path in new.items():
        if name in orig:continue
        tree=ET.fromstring(dst.read(path))
        for c in tree.findall(f'.//{{{NS}}}sheetData/{{{NS}}}row/{{{NS}}}c'):
            assert int(c.attrib.get('s','0'))<style_count
        relpath=posixpath.dirname(path)+'/_rels/'+posixpath.basename(path)+'.rels'
        if relpath in dst.namelist():
            for link in ET.fromstring(dst.read(relpath)):
                target=link.attrib['Target']
                if link.attrib.get('TargetMode')!='External':
                    target=target.lstrip('/') if target.startswith('/') else posixpath.normpath(posixpath.dirname(path)+'/'+target)
                    assert target in dst.namelist(),(name,target)
    assert len([x for x in dst.namelist() if re.fullmatch(r'xl/charts/chart\d+\.xml',x)])==2

# Recompute the two populations directly, without importing the builder.
raw=pd.read_excel(INPUT,sheet_name=['原始SQP合并','广告词搜索数据'])
s=raw['原始SQP合并'];a=raw['广告词搜索数据'].iloc[:,:26].copy()
s['key']=s['搜索查询'].map(norm);a['key']=a['客户搜索词'].map(norm)
expected_s=s.groupby('key')[['曝光量：ASIN 计数','点击量：ASIN 数量','加购：ASIN 数量','下单成交：ASIN 数量']].sum()
expected_a=a.groupby('key')[['展示量','点击量','花费','7天总销售额','7天总订单数(#)','7天总销售量(#)']].sum()
q=pd.read_csv(OUT/'一_SQP全词.csv').set_index('搜索词')
ad=pd.read_csv(OUT/'二_广告全词.csv').set_index('搜索词')
co=pd.read_csv(OUT/'三_共同词.csv').set_index('搜索词')
assert set(q.index)==set(expected_s.index)
assert set(ad.index)==set(expected_a.index)
assert set(co.index)==set(q.index)&set(ad.index)
assert not any('广告' in c for c in q.columns)
assert not any(any(x in c for x in ('SQP','样本','共同','覆盖','市场')) for c in ad.columns)
for target,source in [('样本曝光','曝光量：ASIN 计数'),('样本点击','点击量：ASIN 数量'),
                      ('样本加购','加购：ASIN 数量'),('样本购买','下单成交：ASIN 数量')]:
    assert np.allclose(q.loc[expected_s.index,target],expected_s[source])
for target,source in [('广告曝光','展示量'),('广告点击','点击量'),('广告花费','花费'),
                      ('广告销售额','7天总销售额'),('广告订单','7天总订单数(#)'),('广告销售件数','7天总销售量(#)')]:
    assert np.allclose(ad.loc[expected_a.index,target],expected_a[source])
small=q['样本点击'].lt(100)
assert q.loc[small,['点击吸引力指数','购买承接指数']].isna().all().all()
assert q.loc['hal','样本曝光']==q.loc['hal','样本点击']==1
assert co.index.nunique()==1802

wb=openpyxl.load_workbook(BOOK,read_only=True,data_only=True)
for name,expected in [('一_SQP全词',q),('二_广告全词',ad),('三_共同词',co)]:
    ws=wb[name];rows=ws.iter_rows(values_only=True);heads=next(rows)
    for values in rows:
        record=dict(zip(heads,values));query=record['搜索词']
        for col in expected.columns:
            actual=record[col];want=expected.loc[query,col]
            if pd.isna(want):assert actual in (None,''),(name,query,col,actual,want)
            elif isinstance(want,(float,int,np.number)):
                assert actual is not None and math.isclose(actual,float(want),rel_tol=1e-10,abs_tol=1e-10),(name,query,col,actual,want)
            else:assert str(actual)==str(want),(name,query,col,actual,want)
    assert ws.max_row==len(expected)+1
# Corrected brief follows the requested layout and uses orders separately from units.
ws=wb['词简表_校正版'];rows=ws.iter_rows(values_only=True);heads=next(rows)
for values in rows:
    r=dict(zip(heads,values));key=r['搜索词']
    if key in ad.index:
        clicks=ad.loc[key,'广告点击'];spend=ad.loc[key,'广告花费']
        want=spend/clicks if clicks else None
        if want is None:assert r['CPC'] in (None,'')
        else:assert math.isclose(r['CPC'],want,rel_tol=1e-12)
        assert r['广告订单']==ad.loc[key,'广告订单']
        assert r['广告销售件数']==ad.loc[key,'广告销售件数']
    else:assert r['广告点击'] in (None,'') and r['CPC'] in (None,'')
    if r['样本ASIN点击量']<100:assert r['CTR倍率_低样本留空'] in (None,'') and r['购买比倍率_低样本留空'] in (None,'')
wb.close()
f=openpyxl.load_workbook(BOOK,read_only=True,data_only=False)
checks={'CPC':('广告花费','广告点击'),'ACOS':('广告花费','广告销售额'),
        '广告订单CVR':('广告订单','广告点击')}
for name in ('二_广告全词','三_共同词','词简表_校正版'):
    heads=[c.value for c in f[name][1]]
    for metric,(n,d) in checks.items():
        cell=f[name].cell(2,heads.index(metric)+1)
        nr=openpyxl.utils.get_column_letter(heads.index(n)+1)+'2'
        dr=openpyxl.utils.get_column_letter(heads.index(d)+1)+'2'
        assert cell.data_type=='f' and f'{nr}/{dr}' in cell.value,(name,metric,cell.value)
f.close()
result={'status':'passed','original_sheets_preserved':5,'total_sheets':32,
        'sqp_terms':len(q),'ad_terms':len(ad),'common_terms':len(co),
        'scope_separation':'passed','source_totals':'passed','formula_caches':'passed',
        'CPC_orders_units':'passed','small_sample_indices':'passed','styles_namespaces_charts':'passed',
        'source_sha256':metrics['source_sha256']}
(OUT/'verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2))
print(json.dumps(result,ensure_ascii=False))
