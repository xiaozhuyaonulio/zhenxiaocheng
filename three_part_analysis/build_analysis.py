"""Build three independent research sections on the latest user-supplied workbook.

The uploaded five worksheets are preserved verbatim. New analysis is recomputed
from the two source worksheets; no previous report is used as a numeric input.
Run from the repository root: python three_part_analysis/build_analysis.py
"""
from pathlib import Path
from copy import deepcopy
from collections import Counter
import hashlib, json, math, re, zipfile, posixpath
from lxml import etree as ET
import numpy as np
import pandas as pd
import openpyxl
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import DataBarRule

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'three_part_analysis'
INPUT = ROOT / 'reference/26.9汇总表_三块分析底稿.xlsx'
BOOK = OUT / '26.9汇总表_三块分析版.xlsx'
OUT.mkdir(exist_ok=True)
SCREEN = 100

def norm(x):
    return re.sub(r'\s+', ' ', str(x).strip().lower())

def clean(x):
    if x is None or (not isinstance(x, str) and pd.isna(x)):
        return None
    return x.item() if isinstance(x, np.generic) else x

def div(n, d):
    return n / d if pd.notna(n) and pd.notna(d) and d > 0 else np.nan

def ratio(df, name, n, d):
    df[name] = df[n] / df[d].replace(0, np.nan)

def frac(n, d, clicks=None):
    gate = f'ISNUMBER({clicks}),{clicks}>={SCREEN},' if clicks else ''
    return f'=IFERROR(IF(AND({gate}ISNUMBER({n}),ISNUMBER({d}),{d}>0),{n}/{d},""),"")'

def pp(n, d):
    return f'=IFERROR(IF(AND(ISNUMBER({n}),ISNUMBER({d})),({n}-{d})*100,""),"")'

def pct(v): return f'{v:.2%}'
def money(v): return f'${v:,.2f}'
def count(v): return f'{v:,.0f}'

# Mutually exclusive primary categories, with separate overlapping intent tags.
# These are transparent text rules, not verified product relevance judgments.
CATEGORIES = [
    ('吸管或配件', r'\b(straws?|spouts?|accessor(?:y|ies)|replacement|lids?)\b',
     r'\b(jars?|pitchers?|carafes?|bottles?|cups?|bowls?|containers?)\b'),
    ('碗或备餐', r'\b(bowls?|meal prep)\b', None),
    ('罐类', r'\b(jars?|mason)\b', None),
    ('壶或倒水容器', r'\b(pitchers?|carafes?|jugs?|decanters?)\b', None),
    ('杯或水瓶', r'\b(cups?|tumblers?|bottles?|mugs?)\b', None),
    ('保鲜或储存容器', r'\b(containers?|storage|canisters?|preserv(?:e|ing))\b', None),
    ('咖啡或茶相关', r'\b(coffee|cold brew|tea|brewers?)\b', None),
]
TAGS = [
    ('冷萃', r'\bcold brew\b'), ('茶饮', r'\b(tea|iced tea|sun tea)\b'),
    ('冰箱', r'\b(fridge|refrigerator)\b'), ('户外', r'\b(outside|outdoor|camping)\b'),
    ('储存', r'\b(storage|airtight|preserv(?:e|ing))\b'),
    ('备餐', r'\b(meal prep|lunch)\b'), ('把手', r'\bhandles?\b'),
    ('带盖', r'\blids?\b'), ('玻璃', r'\bglass\b'), ('塑料', r'\bplastic\b'),
    ('容量规格', r'\b\d+(?:\.\d+)?\s*(?:oz|ounce|ounces|ml|liter|litre|quart|gallon)\b'),
]

def classify(q):
    if re.fullmatch(r'b0[a-z0-9]{8}', q):
        return 'ASIN编码'
    for name, pattern, exclude in CATEGORIES:
        if re.search(pattern, q) and (not exclude or not re.search(exclude, q)):
            return name
    return '其他文字词'

def tags(q):
    return '；'.join(name for name, pattern in TAGS if re.search(pattern, q))

def decorate(df):
    df['规则主类别'] = df['搜索词'].map(classify)
    df['场景或属性标签'] = df['搜索词'].map(tags)
    return df

source_hash = hashlib.sha256(INPUT.read_bytes()).hexdigest()
raw = pd.read_excel(INPUT, sheet_name=['原始SQP合并', '广告词搜索数据'])
sraw = raw['原始SQP合并']
araw = raw['广告词搜索数据'].iloc[:, :26]
assert len(sraw) == 3117 and len(araw) == 61169
assert not sraw.duplicated(['ASIN', '搜索查询']).any()
smap = {'搜索查询':'搜索词', '搜索查询量':'市场搜索查询量', '曝光：曝光总量':'市场曝光',
        '点击量：总次数':'市场点击', '购物车添加：总数':'市场加购', '购买：下单总数':'市场购买',
        '曝光量：ASIN 计数':'样本曝光', '点击量：ASIN 数量':'样本点击',
        '加购：ASIN 数量':'样本加购', '下单成交：ASIN 数量':'样本购买'}
s = sraw.rename(columns=smap).copy()
s['搜索词'] = s['搜索词'].map(norm)
s['原稿SQP行'] = np.arange(2, len(s) + 2)
market_cols = ['市场搜索查询量','市场曝光','市场点击','市场加购','市场购买']
own_cols = ['样本曝光','样本点击','样本加购','样本购买']
assert not s.groupby('搜索词')[market_cols].nunique(dropna=False).gt(1).any().any()
market = s.drop_duplicates('搜索词').set_index('搜索词')[market_cols]
q = s.groupby('搜索词')[own_cols].sum().join(market).reset_index()
q = q.merge(s.groupby('搜索词').agg(收录ASIN数=('ASIN','nunique'),
               查询得分最小值=('搜索查询得分','min'),查询得分最大值=('搜索查询得分','max'),
               原稿SQP首行=('原稿SQP行','min')), on='搜索词', validate='one_to_one')
SQP_RATIOS = {
    '样本曝光CTR':('样本点击','样本曝光'), '市场曝光CTR':('市场点击','市场曝光'),
    '样本点击后加购比':('样本加购','样本点击'), '市场点击后加购比':('市场加购','市场点击'),
    '样本点击后购买比':('样本购买','样本点击'), '市场点击后购买比':('市场购买','市场点击'),
    '样本曝光份额':('样本曝光','市场曝光'), '样本点击份额':('样本点击','市场点击'),
    '样本加购份额':('样本加购','市场加购'), '样本购买份额':('样本购买','市场购买'),
    '市场查询点击比':('市场点击','市场搜索查询量'),
    '市场查询加购比':('市场加购','市场搜索查询量'),
    '市场查询购买比':('市场购买','市场搜索查询量'),
    '平均每查询市场曝光':('市场曝光','市场搜索查询量'),
}
for name, (n,d) in SQP_RATIOS.items(): ratio(q, name, n, d)
q['CTR差_百分点'] = (q['样本曝光CTR'] - q['市场曝光CTR']) * 100
q['购买比差_百分点'] = (q['样本点击后购买比'] - q['市场点击后购买比']) * 100
q['同词市场率参考购买计数'] = q['样本点击'] * q['市场点击后购买比']
for name,n,d in [('点击吸引力指数','样本曝光CTR','市场曝光CTR'),
                 ('购买承接指数','样本点击后购买比','市场点击后购买比')]:
    ratio(q,name,n,d)
    q.loc[q['样本点击'].lt(SCREEN),name] = np.nan
q['样本条件'] = np.where(q['样本点击'].ge(SCREEN), '≥100点击；仍需核查', '<100点击；仅记录')

def sqp_hint(r):
    if r['样本点击'] < SCREEN: return '样本点击少，暂不判断'
    if r['市场点击'] < SCREEN or not r['市场曝光CTR'] or pd.isna(r['市场点击后购买比']):
        return '市场基准不足或无定义'
    if r['同词市场率参考购买计数'] < 10: return '购买基准计数少，先观察'
    ctr = 'CTR低于市场' if r['CTR差_百分点'] < 0 else 'CTR不低于市场'
    buy = '购买比低于市场' if r['购买比差_百分点'] < 0 else '购买比不低于市场'
    return ctr + '；' + buy + '（描述）'

decorate(q)
q['描述性线索'] = q.apply(sqp_hint,axis=1)
q = q.sort_values('样本点击',ascending=False).reset_index(drop=True)
q = q[['搜索词','规则主类别','场景或属性标签','市场搜索查询量','收录ASIN数',
       '市场曝光','样本曝光','市场点击','样本点击','市场加购','样本加购','市场购买','样本购买',
       *SQP_RATIOS.keys(),'CTR差_百分点','购买比差_百分点','同词市场率参考购买计数',
       '点击吸引力指数','购买承接指数','样本条件','描述性线索',
       '查询得分最小值','查询得分最大值','原稿SQP首行']]

mapping = pd.read_excel(ROOT/'source_material/raw_data/对应关系.xlsx').dropna(subset=['ASIN'])
assert not mapping['ASIN'].duplicated().any()
sd = s.merge(mapping[['ASIN','品名','颜色','规格']],on='ASIN',how='left',validate='many_to_one')
assert sd['品名'].notna().all()
sd['同词市场率参考购买计数'] = sd['样本点击'] * sd['市场购买'] / sd['市场点击'].replace(0,np.nan)
asin = sd.groupby(['ASIN','品名','颜色','规格'],dropna=False).agg(
    收录词数=('搜索词','size'),样本曝光=('样本曝光','sum'),样本点击=('样本点击','sum'),
    样本加购=('样本加购','sum'),样本购买=('样本购买','sum'),
    同词市场率参考购买计数=('同词市场率参考购买计数','sum')).reset_index()
for name,n,d in [('样本曝光CTR','样本点击','样本曝光'),('样本点击后加购比','样本加购','样本点击'),
                 ('样本点击后购买比','样本购买','样本点击')]: ratio(asin,name,n,d)
asin['点击贡献占比'] = asin['样本点击']/asin['样本点击'].sum()
asin['购买贡献占比'] = asin['样本购买']/asin['样本购买'].sum()
asin['购买比结构参考差_百分点'] = (asin['样本点击后购买比']-
    asin['同词市场率参考购买计数']/asin['样本点击'].replace(0,np.nan))*100
asin['范围提醒'] = '仅各ASIN已导出词；收录词与截取范围不同，不代表全商品业绩'
asin = asin.sort_values('样本点击',ascending=False)
price_cols = ['点击量：价格(中位数)','点击量：ASIN 售价（中位数）','加入购物车：售价（中位数）',
              '加入购物车：ASIN 价格（中位数）','购买次数：价格（中位数）','购买：ASIN 价格（中位数）']
sd['价格核查'] = np.where(sd[price_cols].gt(1000).any(axis=1),'原始价格>1000，待核查','')
detail = sd[['ASIN','品名','颜色','规格','搜索词','搜索查询得分',*own_cols,*price_cols,'价格核查','原稿SQP行']].copy()
for name,n,d in [('样本曝光CTR','样本点击','样本曝光'),('样本点击后加购比','样本加购','样本点击'),
                 ('样本点击后购买比','样本购买','样本点击')]: ratio(detail,name,n,d)

aq = araw.rename(columns={'客户搜索词':'搜索词','展示量':'广告曝光','点击量':'广告点击','花费':'广告花费',
    '7天总销售额':'广告销售额','7天总订单数(#)':'广告订单','7天总销售量(#)':'广告销售件数',
    '7天内广告SKU销售额':'广告SKU销售额','7天内其他SKU销售额':'其他SKU销售额'}).copy()
aq['原始客户词'] = aq['搜索词']
aq['搜索词'] = aq['搜索词'].map(norm)
aq['原稿广告行'] = np.arange(2,len(aq)+2)
aq['客户词类型'] = np.where(aq['搜索词'].str.fullmatch(r'b0[a-z0-9]{8}'),'ASIN编码','文字词')
auto = {'close-match':'自动-紧密','loose-match':'自动-宽泛','substitutes':'自动-同类','complements':'自动-关联'}
def target_type(r):
    if r['投放'] in auto: return auto[r['投放']]
    if r['匹配类型'] in ('EXACT','PHRASE','BROAD'):
        return {'EXACT':'手动-精准','PHRASE':'手动-词组','BROAD':'手动-广泛'}[r['匹配类型']]
    if str(r['投放']).startswith('asin'): return '商品定向-ASIN'
    if str(r['投放']).startswith('category'): return '商品定向-类目'
    if str(r['投放']).startswith('keyword-group'): return '关键词组定向'
    return '其他-待核查'
aq['投放类型'] = aq.apply(target_type,axis=1)
decorate(aq)
adcols = ['广告曝光','广告点击','广告花费','广告销售额','广告订单','广告销售件数','广告SKU销售额','其他SKU销售额']
AD_RATIOS = {'广告CTR':('广告点击','广告曝光'),'广告订单CVR':('广告订单','广告点击'),
             'CPC':('广告花费','广告点击'),'ACOS':('广告花费','广告销售额'),
             'ROAS':('广告销售额','广告花费'),'每订单归因销售额':('广告销售额','广告订单'),
             '其他SKU收入占比':('其他SKU销售额','广告销售额')}
def ad_rates(d):
    for name,(n,de) in AD_RATIOS.items(): ratio(d,name,n,de)
    return d
def ad_agg(keys):
    d = aq.groupby(keys,dropna=False)[adcols].sum()
    trace = aq.groupby(keys,dropna=False).agg(原始记录数=('搜索词','size'),
                原稿广告首行=('原稿广告行','min'),原稿广告末行=('原稿广告行','max'))
    d = ad_rates(d.join(trace).reset_index())
    return d.sort_values('广告花费',ascending=False).reset_index(drop=True)
ad = decorate(ad_agg(['搜索词','客户词类型']))

def ad_hint(r):
    if r['广告订单']==0 and r['广告点击']>=20 and r['广告花费']>=20: return '零单成本复核'
    if r['广告订单']>=5 and r['广告花费']>=50 and r['ACOS']>.4: return '有单高成本复核'
    if r['广告订单']>=10 and r['ACOS']<=.25: return '增长测试候选'
    return '观察或常规维护'
ad['广告筛选线索'] = ad.apply(ad_hint,axis=1)
ad['花费贡献占比'] = ad['广告花费']/ad['广告花费'].sum()
ad['销售贡献占比'] = ad['广告销售额']/ad['广告销售额'].sum()
ad = ad[['搜索词','客户词类型','规则主类别','场景或属性标签',*adcols,*AD_RATIOS.keys(),
         '花费贡献占比','销售贡献占比','广告筛选线索','原始记录数','原稿广告首行','原稿广告末行']]
portfolio = ad_agg(['广告组合名称'])
types = ad_agg(['投放类型'])
pq = decorate(ad_agg(['广告组合名称','搜索词','客户词类型']))
pq['广告筛选线索'] = pq.apply(ad_hint,axis=1)
units = ad_agg(['广告组合名称','广告活动名称','广告组名称','投放','匹配类型','投放类型'])
units['广告筛选线索'] = units.apply(ad_hint,axis=1)
cost_candidates = pq[pq['广告筛选线索'].isin(['零单成本复核','有单高成本复核'])].copy()
growth_candidates = pq[pq['广告筛选线索'].eq('增长测试候选')].copy()
supported_pq = pq[pq['广告点击'].ge(20)].copy()
multi = supported_pq.groupby('搜索词')['广告组合名称'].nunique()
differences = supported_pq[supported_pq['搜索词'].isin(multi[multi.ge(2)].index)].sort_values(
    ['搜索词','广告花费'],ascending=[True,False])

def sqp_group(key):
    g = q.groupby(key)[market_cols+own_cols].sum()
    g = g.join(q.groupby(key)['搜索词'].size().rename('词数')).reset_index()
    for name,(n,d) in SQP_RATIOS.items(): ratio(g,name,n,d)
    g['样本点击贡献占比'] = g['样本点击']/q['样本点击'].sum()
    g['样本购买贡献占比'] = g['样本购买']/q['样本购买'].sum()
    return g.sort_values('样本点击',ascending=False)
sqp_groups = sqp_group('规则主类别')
ad_groups = ad_agg(['客户词类型','规则主类别'])
ad_groups['词数'] = ad_groups.apply(lambda r:ad[(ad['客户词类型']==r['客户词类型']) &
                (ad['规则主类别']==r['规则主类别'])].shape[0],axis=1)

concentration=[]
for key,label,total in [('市场搜索查询量','市场搜索量',q['市场搜索查询量'].sum()),
                        ('样本点击','样本点击',q['样本点击'].sum()),('样本购买','样本购买',q['样本购买'].sum())]:
    ranked=q.sort_values(key,ascending=False)
    for k in (10,50,100):
        v=ranked.head(k)[key].sum()
        concentration.append({'排序指标':label,'头部词数':k,'头部计数':v,'全部收录词计数':total,'头部贡献占比':div(v,total)})
concentration=pd.DataFrame(concentration)
concentration['头部范围']=concentration['排序指标']+'前'+concentration['头部词数'].astype(str)+'词'
ad_concentration=[]
for key in ('广告花费','广告销售额','广告订单'):
    ranked=ad.sort_values(key,ascending=False)
    for k in (10,50,100,1000):
        v=ranked.head(k)[key].sum()
        ad_concentration.append({'排序指标':key,'头部词数':k,'头部计数':v,
                  '全部客户词计数':ad[key].sum(),'头部贡献占比':div(v,ad[key].sum())})
ad_concentration=pd.DataFrame(ad_concentration)
ad_concentration['头部范围']=ad_concentration['排序指标']+'前'+ad_concentration['头部词数'].astype(str)+'词'

S,A=set(q['搜索词']),set(ad['搜索词'])
coverage=[]
for label,terms in [('两表共同收录',S&A),('仅SQP收录',S-A),('仅广告收录',A-S)]:
    sq=q[q['搜索词'].isin(terms)]; ag=ad[ad['搜索词'].isin(terms)]
    coverage.append({'集合':label,'条目数':len(terms),'SQP样本点击':sq['样本点击'].sum() if len(sq) else None,
        'SQP样本购买':sq['样本购买'].sum() if len(sq) else None,
        **{c:ag[c].sum() if len(ag) else None for c in adcols},
        '占全部广告点击':div(ag['广告点击'].sum(),ad['广告点击'].sum()) if len(ag) else None,
        '占全部广告花费':div(ag['广告花费'].sum(),ad['广告花费'].sum()) if len(ag) else None})
coverage=pd.DataFrame(coverage)
common=q.merge(ad.drop(columns=['规则主类别','场景或属性标签']),on='搜索词',how='inner',validate='one_to_one')
assert len(q)==2410 and len(ad)==43846 and len(common)==1802

def cross_hint(r):
    if r['样本点击']<SCREEN or r['同词市场率参考购买计数']<10:
        return 'SQP支持不足；广告独立复核'
    weak = r['购买比差_百分点']<0
    h=r['广告筛选线索']
    if h=='增长测试候选':
        return '广告较低成本但SQP购买比偏低；先核查范围' if weak else '购买比不低且广告达到测试筛选；查利润与商品'
    if h in ('零单成本复核','有单高成本复核'):
        return '购买比偏低且广告待复核；查词意图与商品' if weak else 'SQP购买比不低，广告待复核；查投放成本'
    return '记录同词表现；暂不形成调整结论'
common['交叉验证方向']=common.apply(cross_hint,axis=1)
common['范围核查']='同词不代表同商品或同归因；推广ASIN未对齐，不相减'
common=common.sort_values('广告花费',ascending=False)
cross_candidates=common[common['交叉验证方向'].ne('记录同词表现；暂不形成调整结论') &
                        common['样本点击'].ge(SCREEN)].copy()
sqp_candidates=q[q['样本点击'].ge(SCREEN)].copy()

# A corrected counterpart of the user's 27-column brief, with the same first 27 positions.
brief=pd.DataFrame({
    '搜索词':q['搜索词'],'市场搜索查询量':q['市场搜索查询量'],'收录ASIN数':q['收录ASIN数'],
    '样本ASIN曝光量':q['样本曝光'],'市场曝光量':q['市场曝光'],
    '样本ASIN点击量':q['样本点击'],'市场点击量':q['市场点击'],
    '样本ASIN加购量':q['样本加购'],'市场加购量':q['市场加购'],
    '样本ASIN购买量':q['样本购买'],'市场购买量':q['市场购买'],
    '样本CTR':q['样本曝光CTR'],'市场CTR':q['市场曝光CTR'],
    '样本点击后加购比':q['样本点击后加购比'],'市场点击后加购比':q['市场点击后加购比'],
    '样本点击后购买比':q['样本点击后购买比'],'市场点击后购买比':q['市场点击后购买比'],
    'CTR倍率_低样本留空':q['点击吸引力指数'],'购买比倍率_低样本留空':q['购买承接指数'],
    '样本购买份额':q['样本购买份额']})
lookup=ad.set_index('搜索词')
for name in ['广告点击','广告销售件数','广告花费','ACOS','ROAS','CPC']:
    brief[name]=brief['搜索词'].map(lookup[name])
brief['是否共同收录']=np.where(brief['搜索词'].isin(A),'是','否')
for name in ['广告订单','广告销售额','广告曝光','广告CTR','广告订单CVR','其他SKU收入占比']:
    brief[name]=brief['搜索词'].map(lookup[name])
brief['样本条件']=q['样本条件']
corrected=q.merge(ad.drop(columns=['规则主类别','场景或属性标签']),on='搜索词',how='left',validate='one_to_one')

cached=openpyxl.load_workbook(INPUT,read_only=True,data_only=True)
original_rows=list(cached['总表'].iter_rows(min_row=3,values_only=True))
original_rows=[r for r in original_rows if r[0] is not None]
assert len(original_rows)==len(q)
ql=q.set_index('搜索词'); mismatches=Counter(); errors=Counter(); negatives=Counter()
for r in original_rows:
    expected=ql.loc[norm(r[0])]
    for col,name in [(3,'市场搜索查询量'),(4,'市场曝光'),(5,'样本曝光'),(9,'市场点击'),
                     (12,'样本点击'),(20,'市场加购'),(24,'样本加购'),(33,'市场购买'),(37,'样本购买')]:
        if r[col]!=expected[name]: mismatches[name]+=1
    for col in (57,58,59):
        if isinstance(r[col],(int,float)) and r[col]<0: negatives[get_column_letter(col+1)]+=1
for title in ('总表','广告词简表'):
    for row in cached[title].iter_rows():
        for cell in row:
            if cell.data_type=='e':errors[title]+=1
cached.close()
assert not mismatches,dict(mismatches)
assert np.allclose(aq['广告SKU销售额']+aq['其他SKU销售额'],aq['广告销售额'],atol=.011)
assert q[own_cols].sum().to_dict()==s[own_cols].sum().to_dict()
assert np.allclose(ad[adcols].sum(),aq[adcols].sum())
assert np.allclose(pq[adcols].sum(),aq[adcols].sum())

tot=ad[adcols].sum(); tot_rates={name:div(tot[n],tot[d]) for name,(n,d) in AD_RATIOS.items()}
stot=q[own_cols].sum(); mtot=q[market_cols].sum()
qtop=q.head(1).iloc[0]; atop=asin.iloc[0]
co=coverage.set_index('集合').loc['两表共同收录']
zero=cost_candidates[cost_candidates['广告筛选线索'].eq('零单成本复核')]
high=cost_candidates[cost_candidates['广告筛选线索'].eq('有单高成本复核')]
multi_spend=aq.loc[pd.to_datetime(aq['开始日期'])!=pd.to_datetime(aq['结束日期']),'广告花费'].sum()
stats={'source_file':str(INPUT.relative_to(ROOT)),'source_sha256':source_hash,
       'source_sqp_rows':len(s),'source_ad_rows':len(aq),'sqp_asins':s['ASIN'].nunique(),
       'sqp_queries':len(q),'ad_terms':len(ad),'common_terms':len(common),
       'sqp_only_terms':len(S-A),'ad_only_terms':len(A-S),
       'sqp_totals':stot.to_dict(),'market_deduplicated_totals':mtot.to_dict(),
       'ad_totals':tot.to_dict(),'ad_rates':tot_rates,'coverage':coverage.to_dict('records'),
       'original_formula_errors':dict(errors),'original_negative_differences':dict(negatives),
       'sqp_click_screen_terms':len(sqp_candidates),'cost_candidates':len(cost_candidates),
       'growth_candidates':len(growth_candidates),'cross_candidates':len(cross_candidates),
       'spend_in_multiday_rows':float(multi_spend)}
def safe_json(x):
    if isinstance(x,dict):return {k:safe_json(v) for k,v in x.items()}
    if isinstance(x,list):return [safe_json(v) for v in x]
    return clean(x)
(OUT/'analysis_metrics.json').write_text(json.dumps(safe_json(stats),ensure_ascii=False,indent=2,allow_nan=False))

def findings(rows):return pd.DataFrame(rows,columns=['证据层级','分析问题','数据事实或结果','下一步与边界'])
sqp_findings=findings([
 ['已核对事实','这块数据有多大',f'51个ASIN、3,117条ASIN×词记录、2,410个去重词；样本曝光{count(stot["样本曝光"])}、点击{count(stot["样本点击"])}、加购{count(stot["样本加购"])}、购买{count(stot["样本购买"])}。','只使用SQP源页。每ASIN最多100词；不是51个ASIN的完整流量。'],
 ['已核对事实','核心计数是否可信','逐词核对新版总表的9个核心市场与样本计数字段，全部一致。','保留市场同词取一次、样本计数求和的思路。'],
 ['已核对事实','流量集中在哪些词',f'点击最多的词为{qtop["搜索词"]}，{count(qtop["样本点击"])}点击、{count(qtop["样本购买"])}购买。前10词贡献{pct(q.head(10)["样本点击"].sum()/stot["样本点击"])}样本点击。','查看一_流量集中度与一_SQP全词；头部按各指标分别排序，不混用排名。'],
 ['已核对事实','主力ASIN是谁',f'{atop["ASIN"]}（{atop["品名"]}/{atop["颜色"]}/{atop["规格"]}）贡献{pct(atop["点击贡献占比"])}样本点击、{pct(atop["购买贡献占比"])}样本购买。','点击贡献与购买贡献不同是描述；各ASIN词结构与截取范围不同，不能直接认定商品优劣。'],
 ['诊断线索','哪些词需要检查',f'有{len(sqp_candidates)}个词达到100样本点击展示筛选线，其余{len(q)-len(sqp_candidates)}个词保留计数，不凭倍率判断。','看一_SQP问题候选，按计数与绝对率筛查；市场率参考购买不足10时也标记先观察。筛选线不是显著性标准。'],
 ['已核对事实','极端CTR是否证明优势','hal仅1曝光1点击，CTR100%；同词市场CTR约0.01531%，数学倍率约6,532倍。','两个指数在小样本时留空；保留原始计数。不能据此判断吸引力强。'],
 ['当前不能确定','增长、自然表现与差异原因','只有一个月及截取词样本；没有完整自然/广告拆分、展示位置、页面或库存证据。','不推算自然流量，不预测增长，不把率差写成确定原因。'],
])
for word in ('mason jars','cold brew pitcher','mason jar pitcher'):
    r=ql.loc[word]
    sqp_findings.loc[len(sqp_findings)]=['诊断线索',word,
       f'{count(r["样本点击"])}点击、{count(r["样本购买"])}购买；样本/市场CTR {pct(r["样本曝光CTR"])} / {pct(r["市场曝光CTR"])}；购买/点击 {pct(r["样本点击后购买比"])} / {pct(r["市场点击后购买比"])}。',
       '先检查词意图与商品形态、具体ASIN、价格及页面；数值差异不是已证实的原因或统计显著差异。']
group_top=sqp_groups.iloc[0]
sqp_findings.loc[len(sqp_findings)]=['已核对事实','哪类词贡献较多',
   f'按公开文字规则，“{group_top["规则主类别"]}”共{count(group_top["词数"])}词，贡献{pct(group_top["样本点击贡献占比"])}样本点击、{pct(group_top["样本购买贡献占比"])}样本购买。',
   '分组汇总是观察样本的描述；规则分类需人工复核，不代表已证明商品相关性或整个市场需求。']
ad_findings=findings([
 ['已核对事实','广告整体投入产出',f'广告点击{count(tot["广告点击"])}；花费{money(tot["广告花费"])}；销售额{money(tot["广告销售额"])}；订单{count(tot["广告订单"])}；件数{count(tot["广告销售件数"])}。','单独使用广告源页；订单与件数分开。'],
 ['已核对事实','主要效率指标',f'CTR {pct(tot_rates["广告CTR"])}；CPC {money(tot_rates["CPC"])}；订单CVR {pct(tot_rates["广告订单CVR"])}；ACOS {pct(tot_rates["ACOS"])}；ROAS {tot_rates["ROAS"]:.2f}。','全部按合计分子÷合计分母；不用各行率的简单平均。'],
 ['已核对事实','广告表有多少客户条目',f'61,169条记录，归一化后43,846个客户条目，其中{count(ad["客户词类型"].eq("ASIN编码").sum())}个为ASIN编码。','ASIN编码单独看，不全部称为文字关键词或长尾词。'],
 ['已核对事实','销售构成',f'广告SKU销售额{money(tot["广告SKU销售额"])}；其他SKU销售额{money(tot["其他SKU销售额"])}，占总归因销售{pct(tot_rates["其他SKU收入占比"])}。','缺实际售出SKU与成本，不能以总ACOS直接判断单品利润。'],
 ['待验证建议','哪些花费需要复核',f'组合×词层面：零单成本候选{len(zero)}个，花费{money(zero["广告花费"].sum())}；有单高成本候选{len(high)}个，花费{money(high["广告花费"].sum())}。','这是候选花费规模，不是可节省金额；先核查相关性、具体投放和利润。'],
 ['待验证建议','哪些词可测试增长',f'{len(growth_candidates)}个组合×词满足≥10广告订单且ACOS≤25%的工作筛选线。','25%不是目标或盈亏线。核查利润、库存、商品相关性和预算后再测试。'],
 ['已核对事实','同词投放是否一样',f'有{multi.ge(2).sum()}个客户词在至少两个组合中各有≥20点击，可查看不同组合的成本与订单。','二_同词投放差异按词分组；检查具体投放，不对整店同词统一处理。'],
 ['当前不能确定','每日趋势与利润',f'跨日期区间行承载{pct(multi_spend/tot["广告花费"])}花费，缺可靠每日发生明细；无商品成本。','不按开始日期画每日趋势；不判盈利，不推断广告增量。'],
])
example=pq[pq['搜索词'].eq('mason jar pitcher')].set_index('广告组合名称')
left,right=example.loc['【梅森罐】'],example.loc['【32oz梅森罐】']
ad_findings.loc[len(ad_findings)]=['已核对事实','mason jar pitcher在不同组合',
   f'【梅森罐】{count(left["广告点击"])}点击、{count(left["广告订单"])}订单、ACOS {pct(left["ACOS"])}；【32oz梅森罐】{count(right["广告点击"])}点击、{count(right["广告订单"])}订单、ACOS {pct(right["ACOS"])}。',
   '组合名称不是已验证的ASIN映射；定位具体活动、广告组和投放商品后分别复核。']
ad_findings.loc[len(ad_findings)]=['已核对事实','花费是否集中',
   f'按广告花费排序，前10个客户条目占总花费{pct(ad.head(10)["广告花费"].sum()/tot["广告花费"])}；前100条目占{pct(ad.head(100)["广告花费"].sum()/tot["广告花费"])}。',
   '集中度说明投入分布；不能由此认定预算集中或分散就是合理。查看二_流量集中度。']
cross_findings=findings([
 ['已核对事实','共同词的覆盖范围',f'共同1,802词；仅SQP608词；仅广告42,044条目。共同词含{count(co["SQP样本点击"])}样本点击、{count(co["广告点击"])}广告点击。','共同词只描述两表都收录；仅一表未出现不代表另一来源真实为0。'],
 ['已核对事实','共同词占多少广告',f'共同词占全广告点击{pct(co["占全部广告点击"])}、花费{pct(co["占全部广告花费"])}；花费{money(co["广告花费"])}、销售额{money(co["广告销售额"])}。','这是报告集合的覆盖比例，不是SQP广告占比。第三块不是全广告概览。'],
 ['诊断线索','如何联合筛选',f'{len(cross_candidates)}个共同词达到SQP100点击筛选线并有交叉核查方向；按广告花费排序。','分别保留两边计数、比率和证据；不将广告CVR与SQP购买比视为同一统计口径。'],
 ['待验证建议','两边都支持的候选','SQP购买比不低于同词市场，且广告有一定订单规模与较低ACOS时，列入增长复核。','先对齐推广ASIN、利润和库存；两份相关表现不证明投放造成了增长。'],
 ['待验证建议','两边不一致的词','SQP购买比偏低而广告效益较好，或SQP购买比不低而广告成本较高，列为差异核查。','检查推广商品、其他SKU归因销售、流量构成和投放方式；不直接解释成自然或广告优劣。'],
 ['当前不能确定','能否相减得到自然量','共同词仍未完成推广ASIN、场景、完整流量和购买归因对齐。','不相减、不计算广告占SQP份额、不以正差额当自然量。'],
])
common_lookup=common.set_index('搜索词')
for word,title in [('tea pitcher for fridge','共同词增长复核案例'),
                   ('cold brew coffee maker','共同词成本与意图复核案例'),
                   ('pitcher','两边表现差异案例')]:
    r=common_lookup.loc[word]
    cross_findings.loc[len(cross_findings)]=['诊断线索',title+'：'+word,
       f'SQP {count(r["样本点击"])}点击/{count(r["样本购买"])}购买，购买比{pct(r["样本点击后购买比"])}，同词市场{pct(r["市场点击后购买比"])}；广告{count(r["广告点击"])}点击/{count(r["广告订单"])}订单，花费{money(r["广告花费"])}，ACOS {pct(r["ACOS"])}。',
       r['交叉验证方向']+'。这只是同词并列证据；先对齐推广商品及归因，率差未经显著性检验。']

audit=findings([
 ['已核对事实','输入版本',f'使用本次上传的5页工作簿；SHA256 {source_hash}。','底稿副本另存reference；旧完善版不作为输入。'],
 ['已核对事实','总表核心计数','2,410个词的9个核心计数逐项核对一致。','原稿5页保留，校正版与三块分析重新从源页计算。'],
 ['需修正','广告词简表CPC','原稿Z列用X/V，即ACOS÷广告件数；正确为W/U，即广告花费÷广告点击。','词简表_校正版CPC使用花费/点击；没有点击时留空。'],
 ['需修正','广告订单和件数','原稿总表广告CVR以U列销售件数为分子。','广告订单取源表T列，销售件数取U列；新增订单列与订单CVR。'],
 ['需修正','自然相减列',f'原稿BF/BG/BH有负差值：{dict(negatives)}。即使为正，范围仍不对齐。','新分析不计算自然量；原稿保留页的自然列不用于结论。'],
 ['需修正','分母为0与小样本',f'原稿缓存错误单元格数：{dict(errors)}；hal为1曝光1点击。','新率指标分母0或缺失留空；小样本指数留空，不填假0或截断极端值。'],
 ['需修正','客户词键','统一小写、首尾空格和连续空白。原稿广告存在glass  container双空格。','不进行同义词合并、拼写更正或词干化；共同词按归一化后的精确键匹配。'],
 ['需修正','简表字段标题','原稿CTR、加购率、CVR标题括号列号与当前列位不符；销售量不是订单数。','校正版保留前27列位置，标题改为明确事件口径；补订单、销售额和曝光。'],
 ['范围说明','归因销售一致性','逐广告记录核对广告SKU销售额+其他SKU销售额=总销售额，通过。','只说明报告加总一致，不证明归因销售是广告增量。'],
])

methods=pd.DataFrame([
 ['输入与范围','2026年9月；第一块只读SQP，第二块只读广告，第三块精确取交集。旧报告及豆包结果不作为计算来源。'],
 ['原稿保留','原稿总表、广告词简表、原始SQP合并、广告词搜索数据、Sheet3均保留。原稿自然列、CPC与错误率不用于新结论；优先读分析入口和校正版。'],
 ['市场与样本','同词市场计数一致后取一次；样本计数按收录ASIN求和。样本不是完整品牌、完整商品流量或跨ASIN去重订单。'],
 ['三块独立性','一_SQP全词不含广告字段；二_广告全词包含全广告条目，不以SQP收录过滤；三_共同词仅1,802个交集词。'],
 ['原始与派生','计数、分组、分类标签是输入版本快照。比率、百分点差与小样本指数写入Excel公式和缓存。变更源页后需刷新分组计数。'],
 ['比例加总','组别、ASIN与全量CTR/CVR/ACOS均按合计分子÷合计分母；不平均原始百分比。'],
 ['事件比率','SQP查询点击/加购/购买比、加购/点击、购买/点击是事件计数比，不是去重用户概率，也不是严格逐用户漏斗。'],
 ['百分点','CTR差、购买比差=(样本率-市场率)×100，单位为百分点。不是同比变化；只描述同词绝对率差。'],
 ['小样本指数','样本点击<100或市场率无定义时留空。100只是展示筛选，不是官方标准或统计显著性；低市场基准放大倍数，不以指数排序。'],
 ['SQP线索筛选','≥100样本点击；同词市场率参考购买计数=样本点击×同词市场购买/点击，若<10则先观察。该参考是比较尺，不是预测、损失订单或可挽回订单。10也是工作筛选线。'],
 ['广告筛选线','组合×词层面：零单≥20点击且花费≥20；有单高成本≥5订单、花费≥50、ACOS>40%；增长≥10订单、ACOS≤25%。均为工作复核线，不是目标、盈亏线或操作指令。'],
 ['商品映射','SQP ASIN对应关系使用项目原始对应关系.xlsx；广告尚未得到完整可靠推广ASIN映射。组合名称不能代替已验证的单品归属。'],
 ['中位数与排名','ASIN各阶段原始价格中位数和查询排名只在明细保留；不能平均中位数得到整体中位数，也不能把ASIN查询得分视作自然排名。'],
 ['规则分类','主类别互斥，由公开优先顺序匹配文字；场景属性标签可重叠。分类是文本规则，不是已验证的商品相关性判断。其他文字词保留，品牌与竞品身份未猜测。'],
 ['日期与归因','广告7天归因；含跨日期区间记录，不按开始日期计算每日趋势。货币为USD。SQP购买与广告订单归因口径不同。'],
 ['自然、利润、因果','不相减推算自然量、不计算广告占SQP份额；缺成本不能判利润，缺实验不能判增量与原因。'],
 ['追溯','原稿SQP行/原稿广告首末行对应原稿源页；首末行号只是定位范围，不代表中间每行均属于该词。'],
 ['复算','项目根目录运行python three_part_analysis/build_analysis.py。修改独立底稿源页后重跑；输出校验将核对计数与原稿核心汇总。'],
],columns=['主题','说明'])
rules=pd.DataFrame([
 {'规则类型':'主类别','优先顺序':i,'标签':name,'匹配表达式':pattern,'排除表达式':exclude or '',
  '说明':'按顺序先命中者归入；未命中归其他文字词；b0+8位字母数字优先归ASIN编码'}
 for i,(name,pattern,exclude) in enumerate(CATEGORIES,1)] + [
 {'规则类型':'属性标签','优先顺序':i,'标签':name,'匹配表达式':pattern,'排除表达式':'',
  '说明':'可多标签；标签不能相加作为互斥类别或相关性证明'} for i,(name,pattern) in enumerate(TAGS,1)])

overview_rows=[
 ['一_SQP结论','① 只看SQP',f'{len(q)}个词、{s.ASIN.nunique()}个ASIN',
  '研究市场需求、样本搜索表现、词意图、ASIN分工；不使用广告表。'],
 ['二_广告结论','② 只看广告',f'{len(ad)}个客户条目、{len(aq):,}条记录',
  '研究广告成本、成交贡献、投放差异、复核与增长测试候选。'],
 ['三_共同词结论','③ 只看共同词',f'{len(common)}个词',
  '把同词搜索表现与广告效益并列，形成差异核查与交叉验证清单。'],
 ['词简表_校正版','你的简表校正版','保留前27列布局，补订单等字段','修正CPC、订单/件数、低样本指数和缺失处理；源计数重算。'],
 ['总表_校正版','你的总表校正版','2,410个词的完整搜索与匹配广告信息','保留市场与样本核心指标，补全广告成本；不计算自然相减。'],
 ['指标公式','每个字段怎么来','比例公式与快照取数规则','比例单元格有Excel公式及缓存；分类和汇总是版本快照。'],
 ['校验与修正','原稿核查','原稿5页完整保留于后部','新表用于分析；原稿有误的CPC、自然相减列不作为结论依据。'],
 ['口径与更新','口径和更新','数据事实→诊断线索→验证动作','阈值都是工作筛选线，不自动等同统计可靠或盈利。'],
]
overview=pd.DataFrame(overview_rows,columns=['点击进入','阅读块','数据范围','用途'])

# Derived fields carry formulas; counts and classification remain reproducible snapshots.
GENERAL_RATIOS={**SQP_RATIOS,**AD_RATIOS,
 '样本CTR':('样本ASIN点击量','样本ASIN曝光量'),
 '市场CTR':('市场点击量','市场曝光量'),
 '样本购买份额':('样本购买','市场购买'),
 '头部贡献占比':('头部计数','全部收录词计数')}
BRIEF_RATIOS={'样本CTR':('样本ASIN点击量','样本ASIN曝光量'),
 '市场CTR':('市场点击量','市场曝光量'),'样本点击后加购比':('样本ASIN加购量','样本ASIN点击量'),
 '市场点击后加购比':('市场加购量','市场点击量'),'样本点击后购买比':('样本ASIN购买量','样本ASIN点击量'),
 '市场点击后购买比':('市场购买量','市场点击量'),'样本购买份额':('样本ASIN购买量','市场购买量'),
 **AD_RATIOS}

def field_formula(h,cols,r):
    def addr(name): return get_column_letter(cols.index(name)+1)+str(r)
    pairs=BRIEF_RATIOS if '样本ASIN曝光量' in cols else GENERAL_RATIOS
    if h=='头部贡献占比' and '全部客户词计数' in cols:
        return frac(addr('头部计数'),addr('全部客户词计数'))
    if h in pairs and all(n in cols for n in pairs[h]):
        n,d=pairs[h];return frac(addr(n),addr(d))
    indices={'点击吸引力指数':('样本曝光CTR','市场曝光CTR','样本点击'),
             '购买承接指数':('样本点击后购买比','市场点击后购买比','样本点击'),
             'CTR倍率_低样本留空':('样本CTR','市场CTR','样本ASIN点击量'),
             '购买比倍率_低样本留空':('样本点击后购买比','市场点击后购买比','样本ASIN点击量')}
    if h in indices:
        n,d,c=indices[h];return frac(addr(n),addr(d),addr(c))
    if h=='CTR差_百分点':return pp(addr('样本曝光CTR'),addr('市场曝光CTR'))
    if h=='购买比差_百分点':return pp(addr('样本点击后购买比'),addr('市场点击后购买比'))
    if h=='同词市场率参考购买计数' and '市场点击后购买比' in cols:
        n,d=addr('样本点击'),addr('市场点击后购买比')
        return f'=IFERROR(IF(AND(ISNUMBER({n}),ISNUMBER({d})),{n}*{d},""),"")'
    if h=='购买比结构参考差_百分点':
        a,b,c=addr('样本点击后购买比'),addr('同词市场率参考购买计数'),addr('样本点击')
        return f'=IFERROR(IF(AND(ISNUMBER({a}),ISNUMBER({b}),ISNUMBER({c}),{c}>0),({a}-{b}/{c})*100,""),"")'
    return None

frames={
 '分析入口':overview,'一_SQP结论':sqp_findings,'一_SQP全词':q,'一_词意图分组':sqp_groups,
 '一_流量集中度':concentration,'一_ASIN表现':asin,'一_ASIN词明细':detail,'一_SQP问题候选':sqp_candidates,
 '二_广告结论':ad_findings,'二_广告全词':ad,'二_广告词分组':ad_groups,'二_流量集中度':ad_concentration,
 '二_广告组合':portfolio,'二_投放方式':types,'二_成本复核':cost_candidates,'二_增长测试':growth_candidates,
 '二_同词投放差异':differences,
 '三_共同词结论':cross_findings,'三_覆盖概览':coverage,'三_共同词':common,'三_交叉验证候选':cross_candidates,
 '词简表_校正版':brief,'总表_校正版':corrected,'校验与修正':audit,'口径与更新':methods,'词分类规则':rules,
}
formula_rows=[]
for name,frame in frames.items():
    cols=list(frame.columns)
    for j,h in enumerate(cols,1):
        f=field_formula(h,cols,2)
        formula_rows.append({'分析页':name,'列':get_column_letter(j),'字段':h,
          '类型':'Excel公式' if f else '原始取数、汇总或规则快照',
          '第2行公式或规则':f or ('同词市场取一次；样本/广告计数按相应分组求和；分类与文字按口径页规则。' if len(frame)>50 else '见本页指标、证据或口径；不是新增计算公式。'),
          '注意':'分母0或非数字留空；指数含100点击筛选；百分点差单位为百分点。' if f else '源页修改后需重跑；不平均百分比或中位数。'})
frames['指标公式']=pd.DataFrame(formula_rows)

for name,frame in [('一_SQP全词',q),('一_词意图分组',sqp_groups),('一_ASIN表现',asin),('一_ASIN词明细',detail),
     ('二_广告全词',ad),('二_组合词全量',pq),('二_投放单元全量',units),('二_成本复核',cost_candidates),
     ('二_增长测试',growth_candidates),('三_覆盖概览',coverage),('三_共同词',common),('三_交叉验证候选',cross_candidates)]:
    frame.to_csv(OUT/(name+'.csv'),index=False,encoding='utf-8-sig')

generated=OUT/'_analysis_pages.xlsx'
wb=openpyxl.Workbook(write_only=True)
cache={}; sizes={}
for name,frame in frames.items():
    ws=wb.create_sheet(name)
    color='176C66' if name.startswith('一_') else '205B94' if name.startswith('二_') else '7453A0' if name.startswith('三_') else '173D59'
    ws.sheet_properties.tabColor=color
    ws.freeze_panes='C2' if len(frame.columns)>10 else 'A2'
    ws.sheet_view.showGridLines=False
    ws.row_dimensions[1].height=52
    cols=list(frame.columns)
    for j,h in enumerate(cols,1):
        width=65 if h in ('数据事实或结果','下一步与边界','说明','第2行公式或规则','注意','范围核查') else 40 if h in ('搜索词','广告组合名称','广告活动名称','广告组名称','投放','点击进入','用途','描述性线索','交叉验证方向') else 23
        ws.column_dimensions[get_column_letter(j)].width=width
    head=[]
    for h in cols:
        c=WriteOnlyCell(ws,h);c.data_type='s';c.font=Font(name='微软雅黑',bold=True,color='FFFFFF',size=11)
        c.fill=PatternFill('solid',fgColor=color);c.alignment=Alignment(wrap_text=True,vertical='center');head.append(c)
    ws.append(head)
    cache[name]={}
    for r,record in enumerate(frame.itertuples(index=False,name=None),2):
        if len(frame)<30:ws.row_dimensions[r].height=90 if name not in ('一_流量集中度','二_流量集中度') else 30
        cells=[]
        for j,(h,value) in enumerate(zip(cols,record),1):
            value=clean(value); f=field_formula(h,cols,r)
            c=WriteOnlyCell(ws,f or value)
            if f:cache[name][f'{get_column_letter(j)}{r}']=value
            elif isinstance(value,str):c.data_type='s'
            c.font=Font(name='微软雅黑',size=10)
            c.alignment=Alignment(vertical='top',wrap_text=len(frame)<30 or h in ('描述性线索','交叉验证方向'))
            if isinstance(value,(int,float)) or f:
                if h.endswith('占比') or h.endswith('份额') or h.startswith('占全部') or h in SQP_RATIOS and not h.startswith('平均') or h in ('ACOS','广告CTR','广告订单CVR','样本CTR','市场CTR'):
                    c.number_format='0.00%'
                elif '百分点' in h or '指数' in h or '倍率' in h or h=='ROAS' or h.startswith('平均'):
                    c.number_format='0.00'
                elif any(x in h for x in ('花费','销售额','CPC','价格','售价')):
                    c.number_format='#,##0.00'
                else:c.number_format='#,##0.##' if h=='同词市场率参考购买计数' else '#,##0'
            if name=='分析入口' and j==1:
                c.hyperlink=Hyperlink(ref=f'A{r}',location=f"'{value}'!A1",display=value)
                c.font=Font(name='微软雅黑',color='0563C1',underline='single',size=11)
            cells.append(c)
        ws.append(cells)
    nr,nc=len(frame)+1,len(cols);sizes[name]=(nr,nc)
    if nr>1:ws.auto_filter.ref=f'A1:{get_column_letter(nc)}{nr}'
    for target in ('样本点击','广告花费'):
        if target in cols and nr>2 and nr<10000:
            col=get_column_letter(cols.index(target)+1)
            ws.conditional_formatting.add(f'{col}2:{col}{nr}',DataBarRule(start_type='min',end_type='max',color=color,showValue=True))
    if name in ('一_流量集中度','二_流量集中度'):
        chart=BarChart();chart.type='col';chart.title='头部词贡献占比（各指标独立排序）'
        chart.y_axis.title='贡献占比';chart.y_axis.numFmt='0%';chart.y_axis.scaling.max=1
        chart.add_data(Reference(ws,min_col=5,min_row=1,max_row=nr),titles_from_data=True)
        chart.set_categories(Reference(ws,min_col=6,min_row=2,max_row=nr))
        chart.width=24;chart.height=12;ws.add_chart(chart,'H2')
wb.save(generated)

# Merge new pages into the original OOXML package without reserializing original
# worksheets, source formulas, styles, cached values, or shared strings.
NS='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
REL='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
PKG='http://schemas.openxmlformats.org/package/2006/relationships'
CT='http://schemas.openxmlformats.org/package/2006/content-types'
# lxml retains WPS/Excel namespace declarations referenced by mc:Ignorable.
# ElementTree drops unused declarations and can make the original metadata invalid.
def resolve(target):return target.lstrip('/') if target.startswith('/') else posixpath.normpath('xl/'+target)

def merge_styles(base,new):
    # Copy the new workbook styles, rebasing component indices into the base.
    offsets={};num_map={}
    bn=base.find(f'{{{NS}}}numFmts');nn=new.find(f'{{{NS}}}numFmts')
    if bn is None:bn=ET.Element(f'{{{NS}}}numFmts',{'count':'0'});base.insert(0,bn)
    next_num=max([163]+[int(n.attrib['numFmtId']) for n in bn])+1
    if nn is not None:
        for node in nn:
            old=int(node.attrib['numFmtId']); existing=next((x for x in bn if x.attrib['formatCode']==node.attrib['formatCode']),None)
            if existing is not None:num_map[old]=int(existing.attrib['numFmtId'])
            else:
                cp=deepcopy(node);cp.attrib['numFmtId']=str(next_num);num_map[old]=next_num;next_num+=1;bn.append(cp)
    bn.attrib['count']=str(len(bn))
    for tag in ('fonts','fills','borders','cellStyleXfs','cellXfs','dxfs'):
        b=base.find(f'{{{NS}}}{tag}');n=new.find(f'{{{NS}}}{tag}')
        if n is None:continue
        if b is None:b=ET.Element(f'{{{NS}}}{tag}',{'count':'0'});base.append(b)
        offsets[tag]=len(b)
        for node in n:
            cp=deepcopy(node)
            if tag in ('cellStyleXfs','cellXfs'):
                for attr,component in [('fontId','fonts'),('fillId','fills'),('borderId','borders')]:
                    if attr in cp.attrib:cp.attrib[attr]=str(int(cp.attrib[attr])+offsets[component])
                if 'numFmtId' in cp.attrib:cp.attrib['numFmtId']=str(num_map.get(int(cp.attrib['numFmtId']),int(cp.attrib['numFmtId'])))
                if 'xfId' in cp.attrib:cp.attrib['xfId']=str(int(cp.attrib['xfId'])+offsets['cellStyleXfs'])
            b.append(cp)
        b.attrib['count']=str(len(b))
    return offsets

with zipfile.ZipFile(INPUT) as original,zipfile.ZipFile(generated) as new:
    base_styles=ET.fromstring(original.read('xl/styles.xml'))
    offsets=merge_styles(base_styles,ET.fromstring(new.read('xl/styles.xml')))
    work=ET.fromstring(original.read('xl/workbook.xml'))
    sheets=work.find(f'{{{NS}}}sheets'); original_sheet_nodes=list(sheets)
    rels=ET.fromstring(original.read('xl/_rels/workbook.xml.rels'))
    new_work=ET.fromstring(new.read('xl/workbook.xml'))
    new_rels={x.attrib['Id']:x.attrib['Target'] for x in ET.fromstring(new.read('xl/_rels/workbook.xml.rels'))}
    content_types=ET.fromstring(original.read('[Content_Types].xml'))
    new_ct=ET.fromstring(new.read('[Content_Types].xml'))
    next_id=max(int(x.attrib['Id'][3:]) for x in rels if re.fullmatch('rId[0-9]+',x.attrib['Id']))+1
    next_sheet=max(int(x.attrib['sheetId']) for x in original_sheet_nodes)+1
    replacements={}; paths={}
    for i,node in enumerate(new_work.find(f'{{{NS}}}sheets')):
        name=node.attrib['name'];old_path=resolve(new_rels[node.attrib[f'{{{REL}}}id']])
        new_path=f'xl/worksheets/analysis{i+1}.xml';paths[old_path]=new_path
        cp=deepcopy(node);cp.attrib['sheetId']=str(next_sheet+i);rid=f'rId{next_id+i}';cp.attrib[f'{{{REL}}}id']=rid
        sheets.insert(i,cp)
        rels.append(ET.Element(f'{{{PKG}}}Relationship',{'Id':rid,'Type':REL+'/worksheet','Target':'/'+new_path}))
        xml=ET.fromstring(new.read(old_path))
        nr,nc=sizes[name];xml.insert(1,ET.Element(f'{{{NS}}}dimension',{'ref':f'A1:{get_column_letter(nc)}{nr}'}))
        # Default style for new worksheets must not inherit a nonzero original style.
        for c in xml.findall(f'.//{{{NS}}}sheetData/{{{NS}}}row/{{{NS}}}c'):
            c.attrib['s']=str(int(c.attrib.get('s','0'))+offsets['cellXfs'])
            address=c.attrib['r']
            if address in cache[name]:
                val=cache[name][address];v=c.find(f'{{{NS}}}v')
                if v is None:v=ET.SubElement(c,f'{{{NS}}}v')
                v.text=str(val) if val is not None else None
                if val is None:c.attrib['t']='str';v.text=''
        for col in xml.findall(f'.//{{{NS}}}cols/{{{NS}}}col'):
            if 'style' in col.attrib:col.attrib['style']=str(int(col.attrib['style'])+offsets['cellXfs'])
        for rule in xml.findall(f'.//{{{NS}}}cfRule'):
            if 'dxfId' in rule.attrib:rule.attrib['dxfId']=str(int(rule.attrib['dxfId'])+offsets.get('dxfs',0))
        replacements[new_path]=ET.tostring(xml,encoding='utf-8')
        old_rel=posixpath.dirname(old_path)+'/_rels/'+posixpath.basename(old_path)+'.rels'
        if old_rel in new.namelist():
            new_rel=posixpath.dirname(new_path)+'/_rels/'+posixpath.basename(new_path)+'.rels'
            replacements[new_rel]=new.read(old_rel)
    # Original sheets are unchanged and moved to the end of the workbook.
    for name in work.findall(f'.//{{{NS}}}definedName'):
        if 'localSheetId' in name.attrib:
            name.attrib['localSheetId']=str(int(name.attrib['localSheetId'])+len(frames))
    view=work.find(f'{{{NS}}}bookViews/{{{NS}}}workbookView')
    if view is not None:view.attrib.update({'activeTab':'0','firstSheet':'0'})
    calc=work.find(f'{{{NS}}}calcPr')
    if calc is None:calc=ET.SubElement(work,f'{{{NS}}}calcPr')
    # All new formulas have caches. Avoid forcing a full recalculation of the
    # retained original workbook's expensive whole-column array formulas on open.
    calc.attrib.update({'calcMode':'auto','fullCalcOnLoad':'0','forceFullCalc':'0'})
    app=ET.fromstring(original.read('docProps/app.xml'))
    appns='http://schemas.openxmlformats.org/officeDocument/2006/extended-properties'
    vt='http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes'
    for n in app.findall(f'.//{{{appns}}}HeadingPairs//{{{vt}}}i4'):
        n.text=str(len(sheets))
    parts=app.find(f'{{{appns}}}TitlesOfParts/{{{vt}}}vector')
    if parts is not None:
        for n in list(parts):parts.remove(n)
        parts.attrib['size']=str(len(sheets))
        for node in sheets:ET.SubElement(parts,f'{{{vt}}}lpstr').text=node.attrib['name']
    for node in new_ct:
        if node.tag==f'{{{CT}}}Override':
            old=node.attrib['PartName'].lstrip('/')
            if old in paths:
                cp=deepcopy(node);cp.attrib['PartName']='/'+paths[old];content_types.append(cp)
            elif old.startswith(('xl/charts/','xl/drawings/')):
                assert old not in original.namelist();content_types.append(deepcopy(node))
        elif not any(x.tag==node.tag and x.attrib.get('Extension')==node.attrib.get('Extension') for x in content_types):
            content_types.append(deepcopy(node))
    replacements.update({'xl/workbook.xml':ET.tostring(work,encoding='utf-8'),
       'xl/_rels/workbook.xml.rels':ET.tostring(rels,encoding='utf-8'),
       'xl/styles.xml':ET.tostring(base_styles,encoding='utf-8'),
       'docProps/app.xml':ET.tostring(app,encoding='utf-8'),
       '[Content_Types].xml':ET.tostring(content_types,encoding='utf-8')})
    with zipfile.ZipFile(BOOK,'w',zipfile.ZIP_DEFLATED,compresslevel=6) as dest:
        for item in original.infolist():dest.writestr(item,replacements.pop(item.filename,original.read(item.filename)))
        for path,content in replacements.items():dest.writestr(path,content)
        for path in new.namelist():
            if path.startswith(('xl/charts/','xl/drawings/')):dest.writestr(path,new.read(path))
generated.unlink()
assert hashlib.sha256(INPUT.read_bytes()).hexdigest()==source_hash

# Compact stand-alone report mirrors the workbook, with all figures generated.
report=['# 9月三块数据分析\n',
 '这次按三个独立数据集展开：先只看SQP的2,410词，再只看全量广告43,846条目，最后看共同1,802词。计算直接来自本次上传底稿的两张源数据页。\n']
for title,frame in [('第一块：SQP全部收录词',sqp_findings),('第二块：全量广告客户词',ad_findings),('第三块：两表共同词',cross_findings)]:
    report.append('## '+title+'\n')
    for r in frame.to_dict('records'):
        report.append(f'**{r["分析问题"]}（{r["证据层级"]}）**\n\n{r["数据事实或结果"]}\n\n{r["下一步与边界"]}\n')
report.append('## 怎么读和更新\n\n先打开工作簿的“分析入口”，进入三块结论。各块全词页可筛选并核对计数；你的表格布局保留在“词简表_校正版”，完整指标在“总表_校正版”。原稿5页保留于工作簿后部，其自然相减、错误CPC和件数CVR不用于新结论。\n\n所有筛选线都是工作复核规则，没有目标ACOS和成本时不判利润，不将候选花费当成可节省金额。源数据更新后运行脚本刷新。\n')
(OUT/'三块分析说明与结论.md').write_text('\n'.join(report))
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
doc=Document();style=doc.styles['Normal'];style.font.name='Calibri';style.font.size=Pt(10)
style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'微软雅黑')
for line in '\n'.join(report).splitlines():
    if not line.strip():continue
    if line.startswith('# '):doc.add_heading(line[2:],0)
    elif line.startswith('## '):doc.add_heading(line[3:],1)
    else:doc.add_paragraph(line.replace('**',''))
doc.save(OUT/'三块分析说明与结论.docx')
(OUT/'README.md').write_text('''# 在用户新版汇总表上做的三块分析

先打开 **26.9汇总表_三块分析版.xlsx → 分析入口**。

- 第一块：一_SQP结论、一_SQP全词；全部2,410个词，仅使用SQP。
- 第二块：二_广告结论、二_广告全词；全部43,846个客户条目，包括ASIN编码，仅使用广告源页。
- 第三块：三_共同词结论、三_共同词；精确交集1,802词。
- 词简表_校正版保留用户简表前27列布局；修正CPC、缺失值与小样本指数，补广告订单等字段。
- 总表_校正版补齐搜索与匹配广告指标；不计算自然相减。
- 原稿5个页签完整保留在后部。原稿错误指标不用于新分析。

说明与结论另有Markdown和Word。完整组合词、投放单元也提供CSV；它们不因看板筛选丢失。

输入：reference/26.9汇总表_三块分析底稿.xlsx。本次上传的底稿不修改；只从其两张原始数据页重新计算。商品信息另用source_material/raw_data/对应关系.xlsx。

复算：在项目根目录执行 `python three_part_analysis/build_analysis.py`。校验：`python three_part_analysis/verify_analysis.py`。源页更新后需重跑分组快照；比率已写Excel公式并保留缓存。

分类规则、小样本和广告筛选线见工作簿。筛选线不是官方标准、显著性检验或盈利目标。共同词不代表推广商品和归因范围已对齐，不能相减得到自然量。
''')
print(json.dumps({'workbook':str(BOOK),'new_analysis_sheets':len(frames),'original_sheets_preserved':5,
       'formula_cells':sum(len(v) for v in cache.values()),'workbook_bytes':BOOK.stat().st_size},ensure_ascii=False))
