"""Document existing fields and formulas, without changing business metrics or the data workbook."""
from pathlib import Path
import ast,re
import pandas as pd
import openpyxl
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter

ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'user_summary_review'
wb=openpyxl.load_workbook(OUT/'26.9汇总表_完善版.xlsx',read_only=True,data_only=False)
boardheads=[c.value for c in wb['重点词看板'][1]]
totalheads=[c.value for c in wb['总表'][1]]
assert len(boardheads)==25 and len(totalheads)==71
source=ast.parse((OUT/'review_user_summary.py').read_text())
ratios={}
for node in source.body:
    if isinstance(node,ast.Assign) and any(isinstance(x,ast.Name) and x.id=='ratio_defs' for x in node.targets):
        ratios=ast.literal_eval(node.value)
assert len(ratios)==28
def lookup(col):
    return f'=XLOOKUP($A2,\'总表\'!$A$2:$A$2411,\'总表\'!${col}$2:${col}$2411,"")'
def frac(n,d,clicks=None):
    screen=f'ISNUMBER({clicks}),{clicks}>=100,' if clicks else ''
    return f'=IFERROR(IF(AND({screen}ISNUMBER({n}),ISNUMBER({d}),{d}>0),{n}/{d},""),"")'
marketbuy='XLOOKUP($A2,\'总表\'!$A$2:$A$2411,\'总表\'!$AH$2:$AH$2411,"")'
adsales='XLOOKUP($A2,\'总表\'!$A$2:$A$2411,\'总表\'!$BB$2:$BB$2411,"")'
board_defs={
 1:('文本键','原始搜索词；不计算','', '作为跨表匹配键。'),
 2:('原始取数','同词市场搜索查询量；取一次，不对ASIN求和',lookup('D'),'原报告值。'),
 3:('计数汇总','同词已导出ASIN去重数量',lookup('B'),'不是广告投放商品数，也不是市场竞品数。'),
 4:('计数汇总','Σ同词各收录ASIN曝光',lookup('F'),'仅已导出ASIN×词样本。'),
 5:('计数汇总','Σ同词各收录ASIN点击',lookup('M'),''),
 6:('计数汇总','Σ同词各收录ASIN加购',lookup('Y'),''),
 7:('计数汇总','Σ同词各收录ASIN购买计数',lookup('AL'),'不是跨ASIN去重订单。'),
 8:('比率计算','样本点击÷样本曝光',frac('E2','D2'),'百分比格式。'),
 9:('比率计算/取数','同词市场点击÷市场曝光',lookup('L'),'市场分母未在看板单列；从总表取已算CTR。'),
 10:('比率计算','样本加购÷样本点击',frac('F2','E2'),'次数比，不是去重用户概率。'),
 11:('比率计算/取数','同词市场加购÷市场点击',lookup('X'),'从总表取已算指标。'),
 12:('比率计算','样本购买÷样本点击',frac('G2','E2'),'购买计数/点击。'),
 13:('比率计算/取数','同词市场购买÷市场点击',lookup('AK'),'从总表取已算指标。'),
 14:('有展示筛选的指数','样本CTR÷市场CTR；样本点击<100时留空',frac('H2','I2','E2'),'倍数仅辅助描述；100点击是展示筛选线，不保证稳定。低市场CTR会放大倍数，须看曝光、点击和绝对率。'),
 15:('有展示筛选的指数','样本点击后购买比÷市场点击后购买比；样本点击<100时留空',frac('L2','M2','E2'),'不是显著性检验；仍须核查购买计数和绝对率，不能按倍数排名。'),
 16:('比率计算','样本购买÷同词市场购买',frac('G2',marketbuy),'不能除以市场查询量；该值在总表AM。'),
 17:('计数汇总','Σ同词广告点击',lookup('AZ'),'广告源表M列；按归一化客户词。'),
 18:('计数汇总','Σ同词广告订单',lookup('BF'),'广告源表T列；U列是件数。'),
 19:('计数汇总','Σ同词广告花费',lookup('BA'),'金额USD。'),
 20:('比率计算','同词广告花费÷同词广告7天总销售额',frac('S2',adsales),'销售额不在看板单列，从总表BB按词取。'),
 21:('比率计算','同词广告花费÷广告点击',frac('S2','Q2'),'每次点击金额。'),
 22:('比率计算/取数','Σ同词其他SKU销售额÷Σ同词广告总销售额',lookup('BJ'),'源表Z列/Q列的同词合计；不是其他SKU件数占比。'),
 23:('集合标签','按是否在两份源表出现标记；不计算数值指标',lookup('BO'),'未出现不等于真实0，也不等于纯自然。'),
 24:('规则标签','按样本点击是否≥100标记；不计算业务比率',lookup('BP'),'100是观察阈值，不是官方显著性标准。'),
 25:('诊断标签','依据样本量和两个指数进行描述性分类；不计算新的数值指标',lookup('BQ'),'标签不是确定原因或执行指令。'),
}
records=[]
for i,h in enumerate(boardheads,1):
    typ,math,formula,note=board_defs[i]
    records.append({'列':get_column_letter(i),'字段':h,'类型':typ,'计算口径或取数方式':math,'Excel第2行公式或取数表达式':formula,'说明':note})
board=pd.DataFrame(records)

sqkey="'原始SQP合并'!$B$2:$B$3118"
adkey="'广告词搜索数据'!$AA$2:$AA$61170"
criteria='SUBSTITUTE(SUBSTITUTE(SUBSTITUTE($A2,"~","~~"),"*","~*"),"?","~?")'
def source_sum(sheet,col,key):
    end=3118 if sheet=='原始SQP合并' else 61170
    return f'=SUMIF({key},{criteria},\'{sheet}\'!${col}$2:${col}${end})'
def sq_lookup(col):return f'=XLOOKUP($A2,{sqkey},\'原始SQP合并\'!${col}$2:${col}$3118,"")'
sourcecols={4:'D',5:'E',10:'H',18:'N',19:'O',20:'P',21:'Q',31:'W',32:'X',33:'Y',34:'Z',45:'AF',46:'AG',47:'AH',48:'AI'}
owncols={6:'F',13:'J',25:'S',38:'AB'}
adcols={51:'L',52:'M',53:'P',54:'Q',55:'U',58:'T'}
price_source={16:'L',29:'U',43:'AD'}
repprice={17:('BK','M','点击'),30:('BL','V','加购'),44:('BM','AE','购买')}
total_records=[]
for i,h in enumerate(totalheads,1):
    col=get_column_letter(i);typ='';math='';formula='';note=''
    if i in ratios:
        n,d=ratios[i];typ='比率或指数计算';math=f'{totalheads[n-1]}÷{totalheads[d-1]}'
        formula=frac(get_column_letter(n)+'2',get_column_letter(d)+'2','M2' if i in (49,50) else None)
        note='任一输入非数字、分母0或缺失均留空；不要额外乘100，百分比由单元格格式显示。'
        if i in (49,50):
            typ='有展示筛选的指数';math+='；样本点击<100时留空'
            note='普通小数/倍数；100点击是展示筛选线，不是显著性或稳定性保证。低市场基准放大比值，先看原始计数和绝对率，不按指数大小排序。'
    elif i in owncols:
        typ='计数汇总';math='按同一搜索词，对各已导出ASIN的原始计数求和'
        formula=source_sum('原始SQP合并',owncols[i],sqkey)
        note='数据中ASIN×词不重复；新增重复记录前须去重。'
    elif i in adcols:
        typ='计数汇总';math='按同一归一化客户搜索词，对所有广告记录求和'
        formula=source_sum('广告词搜索数据',adcols[i],adkey)
        note='先在广告源表新增AA归一化客户词辅助列；无该词时应留空，不能只用SUMIF的0判断。'
        formula=f'=IF(COUNTIF({adkey},{criteria})>0,{formula[1:]},"")'
    elif i in sourcecols:
        typ='原始取数';math=f'同词原始SQP合并页{sourcecols[i]}列；一致后取一次，不对ASIN求和'
        formula=sq_lookup(sourcecols[i]);note='如果同词原始市场值不一致，先核查，不能随意取第一条。'
    elif i==1:typ='文本键';math='原始词归一化后去重；不计算数值指标'
    elif i==2:
        typ='去重计数';math='同词ASIN去重数量'
        formula=f'=IFERROR(ROWS(UNIQUE(FILTER(\'原始SQP合并\'!$A$2:$A$3118,{sqkey}=$A2))),0)'
    elif i==3:
        typ='排名描述';math='同词各ASIN查询得分的最小值—最大值；不计算品牌统一排名'
        note='不能平均ASIN查询排名来代表品牌排名；当前显示范围。'
    elif i in price_source:
        typ='原始价格取数';math=f'原始SQP合并页{price_source[i]}列同词市场中位数；不自行重算'
        note='同词市场价格一致才可取一次；空值保持空，原始>1000的可疑值在本版标记待核查。'
    elif i in repprice:
        ac,pc,stage=repprice[i];typ='代表价格取数'
        math=f'读取{ac}列指定ASIN在该词下的{stage}价格中位数（源表{pc}列）'
        note='不是合并中位数；不对各ASIN中位数取平均或加权后仍叫整体中位数。'
    elif i==62:
        typ='比率计算';math='同词广告源表Z列其他SKU销售额合计÷总表BB列总销售额'
        summ=source_sum('广告词搜索数据','Z',adkey)[1:]
        formula=frac(summ,'BB2');note='广告AA辅助列与总表搜索词必须统一。'
    elif i in [63,64,65]:
        stage={63:'点击',64:'加购',65:'购买'}[i];typ='代表ASIN选择规则'
        math=f'同词{stage}计数最多的ASIN；并列按ASIN字典序选第一项；无事件则留空'
        note='属于选择规则，不计算新的业务指标；本版已有选择结果。'
    elif i==66:typ='说明标签';math='检查各ASIN价格是否不同、是否有可疑价格；不计算业务比率'
    elif i==67:typ='集合标签';math='检查词是否出现于SQP与广告报告；不计算数值指标'
    elif i==68:typ='样本标签';math='依据样本点击是否≥100标记；不计算业务比率'
    elif i==69:typ='诊断标签';math='样本量与吸引力/承接指数的描述性规则；不计算新的数值指标'
    elif i==70:typ='口径提示';math='文字说明；不计算'
    else:raise AssertionError((i,h))
    total_records.append({'列':col,'字段':h,'类型':typ,'计算口径或取数方式':math,'Excel第2行公式或取数表达式':formula,'说明':note})
total=pd.DataFrame(total_records)
notes=pd.DataFrame([
 ['适用文件','26.9汇总表_完善版.xlsx；重点词看板A:Y共25列、总表A:BS共71列。公式按第2行示例写，向下填充时行号随之变化。'],
 ['为什么没有公式','部分列是计算快照中的原始值或按词汇总结果，部分是文字规则。指南区分计算、汇总、取数和不计算；不发明未知指标。'],
 ['排序不同','重点词看板与总表的排序不同；从总表取数使用XLOOKUP按A列搜索词匹配，不直接写=总表!D2等相同行号引用。'],
 ['匹配一致','SQP原稿3,117行词与归一化词一致。广告有1行glass  container带双空格，直接按原词SUMIF会漏合并；应先统一空格。'],
 ['广告辅助列','若直接在Excel源表汇总广告，请在广告词搜索数据新增AA列，AA1填归一化客户词，AA2填=LOWER(TRIM(K2))，向下填充至61170行。本次输入只有普通双空格差异。若未来有制表符等空白需额外清洗。'],
 ['无需辅助列的情况','看板可以按词从现有总表取已经归一化计算的结果。使用看板取数公式时，不需要新增广告AA列。'],
 ['数值与格式','CTR、CVR、份额、加购比、购买比、ACOS和收入占比按百分比显示；指数、ROAS按普通小数/倍数，CPC按金额。公式不额外乘100。'],
 ['原始系统百分数字段','原始SQP的32.61代表32.61%；本指南用直接计数相除得到0.3261后设置百分比格式，不把原始32.61再次当0.3261格式化。'],
 ['零与空白','计数有记录且为0，保留0；任一输入不是数字或比例分母为0，留空。指数上游为空也继续为空，空白不是0%。'],
 ['指数展示筛选','看板N/O与总表AW/AX仅在样本点击≥100且上游有效时显示。少于100时保留计数与绝对率，指数留空；100不是统计显著性门槛。市场基准极低仍须核查，不用倍数排名。'],
 ['市场去重','同词市场搜索、曝光、点击、加购、购买值一致后只取一次；ASIN计数可按已收录ASIN求和。'],
 ['价格和排名','市场/ASIN中位数取原始值；跨ASIN整体中位数无法精确还原，不计算。ASIN排名不平均成品牌排名。'],
 ['自然量','理论上同范围完整总量减广告量可分解；本次范围与购买归因未对齐，不计算自然量、广告占比。'],
 ['覆盖与版本','2410个词只覆盖部分广告，范围固定在当前输入版本。新增原始数据时应更新公式范围，并重新核对计数和结论。'],
 ['兼容性','XLOOKUP、FILTER和UNIQUE需要支持这些函数的Excel/WPS版本。指南是公式说明，不会修改主数据工作簿。'],
],columns=['主题','说明'])
target=OUT/'字段计算公式说明.xlsx'
with pd.ExcelWriter(target,engine='openpyxl') as writer:
    for name,df in [('重点词看板25列',board),('总表71列',total),('使用说明',notes)]:
        df.to_excel(writer,sheet_name=name,index=False)
        ws=writer.book[name];ws.freeze_panes='C2';ws.auto_filter.ref=ws.dimensions;ws.sheet_view.showGridLines=False
        for c in ws[1]:c.fill=PatternFill('solid',fgColor='173D59');c.font=Font(name='微软雅黑',bold=True,color='FFFFFF');c.alignment=Alignment(wrap_text=True)
        for j,w in enumerate(([9,29,23,58,88,65] if len(df.columns)==6 else [25,110]),1):ws.column_dimensions[get_column_letter(j)].width=w
        for row in ws.iter_rows(min_row=2):
            for c in row:
                if isinstance(c.value,str):c.data_type='s'
                c.font=Font(name='微软雅黑',size=10);c.alignment=Alignment(vertical='top',wrap_text=True)
            ws.row_dimensions[row[0].row].height=72
def md(df):
    lines=['| '+' | '.join(df.columns)+' |','| '+' | '.join(['---']*len(df.columns))+' |']
    for r in df.itertuples(index=False,name=None):lines.append('| '+' | '.join(str(v).replace('|','/').replace('\n',' ') for v in r)+' |')
    return '\n'.join(lines)
text='# 字段计算公式说明\n\n适用：26.9汇总表_完善版.xlsx。按当前字段列号和第2行列出，向下填充。原始字段不创造计算公式，分类标签不作为新的数值指标。\n\n## 重点词看板：25列\n\n'+md(board)+'\n\n## 完整总表：71列\n\n'+md(total)+'\n\n## 使用说明\n\n'+md(notes)+'\n'
(OUT/'字段计算公式说明.md').write_text(text)
assert len(board)==25 and len(total)==71 and total['列'].iloc[-1]=='BS'
check=openpyxl.load_workbook(target,read_only=True,data_only=False)
assert check['重点词看板25列'].max_row==26 and check['总表71列'].max_row==72
assert all(c.data_type!='f' for s in check for row in s.iter_rows() for c in row)
print('Created literal formula guide:',target,'25+71 fields; documentation does not alter data workbook.')
