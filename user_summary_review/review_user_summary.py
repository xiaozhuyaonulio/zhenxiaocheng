"""Audit and improve the user's September summary using its embedded source sheets.
Run from repository root: python user_summary_review/review_user_summary.py
The original workbook is read-only. Rates use simple Excel formulas with cached values.
"""
from pathlib import Path
from collections import Counter
import hashlib,json,math,re,zipfile,xml.etree.ElementTree as ET
import numpy as np
import pandas as pd
import openpyxl
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'user_summary_review';OUT.mkdir(exist_ok=True)
INPUT=ROOT/'reference/26.9汇总表_原稿.xlsx'
BOOK=OUT/'26.9汇总表_完善版.xlsx'
def norm(v):return re.sub(r'\s+',' ',str(v).strip().lower())
def div(a,b):
    if pd.isna(a) or pd.isna(b) or not b:return np.nan
    return a/b
INDEX_MIN_CLICKS=100 # Display screen only; not a statistical significance threshold.
def excel_ratio(n,d,clicks=None):
    screen=f'ISNUMBER({clicks}),{clicks}>={INDEX_MIN_CLICKS},' if clicks else ''
    return f'=IFERROR(IF(AND({screen}ISNUMBER({n}),ISNUMBER({d}),{d}>0),{n}/{d},""),"")'
def clean(v):
    if v is None or (not isinstance(v,str) and pd.isna(v)):return None
    if isinstance(v,np.generic):return v.item()
    return v

data=pd.read_excel(INPUT,sheet_name=['原始SQP合并','广告词搜索数据'])
sraw,araw=data['原始SQP合并'],data['广告词搜索数据']
assert len(sraw)==3117 and len(araw)==61169
assert not sraw.duplicated(['ASIN','搜索查询']).any()
src_hash=hashlib.sha256(INPUT.read_bytes()).hexdigest()
cached=openpyxl.load_workbook(INPUT,read_only=True,data_only=True)
original=list(cached['总表'].iter_rows(min_row=3,values_only=True))
original=[r for r in original if r[0] is not None]
original_order=[norm(r[0]) for r in original]
original_rows={norm(r[0]):i+3 for i,r in enumerate(original)}
assert len(original_order)==2410 and len(set(original_order))==2410
errors=Counter();negative=Counter()
for row in cached['总表'].iter_rows(min_row=3):
    for c in row:
        if c.data_type=='e':errors[c.column_letter]+=1
        if isinstance(c.value,(int,float)) and c.value<0:negative[c.column_letter]+=1
cached.close()
mp=pd.read_excel(ROOT/'source_material/raw_data/对应关系.xlsx').dropna(subset=['ASIN'])
s=sraw.copy();s['搜索词']=s['搜索查询'].map(norm);s['原稿来源行']=np.arange(2,len(s)+2)
s=s.merge(mp[['ASIN','品名','颜色','规格']],on='ASIN',how='left',validate='many_to_one')
assert s['品名'].notna().all()
assert not s.duplicated(['ASIN','搜索词']).any()
marketcols=['搜索查询量','曝光：曝光总量','点击量：总次数','购物车添加：总数','购买：下单总数']
assert not s.groupby('搜索词')[marketcols].nunique(dropna=False).gt(1).any().any()
first=s.drop_duplicates('搜索词').set_index('搜索词')
groups={q:g for q,g in s.groupby('搜索词',sort=False)}
assert set(original_order)==set(groups)
owncols=['曝光量：ASIN 计数','点击量：ASIN 数量','加购：ASIN 数量','下单成交：ASIN 数量']
for index,col in zip([5,12,24,37],owncols):
    for r in original:
        assert r[index]==groups[norm(r[0])][col].sum(),(r[0],col)

a=araw.rename(columns={'客户搜索词':'原始客户词','展示量':'广告曝光','点击量':'广告点击','花费':'广告花费',
    '7天总销售额':'广告销售额','7天总订单数(#)':'广告订单','7天总销售量(#)':'广告销售件数',
    '7天内广告SKU销售额':'广告SKU销售额','7天内其他SKU销售额':'其他SKU销售额'}).copy()
a['搜索词']=a['原始客户词'].map(norm);a['原稿来源行']=np.arange(2,len(a)+2)
a['客户词类型']=np.where(a['搜索词'].str.fullmatch('b0[a-z0-9]{8}'),'ASIN编码','文本词')
auto={'close-match':'自动-紧密','loose-match':'自动-宽泛','substitutes':'自动-同类','complements':'自动-关联'}
def target(r):
    if r['投放'] in auto:return auto[r['投放']]
    if r['匹配类型'] in ['EXACT','PHRASE','BROAD']:return {'EXACT':'手动-精准','PHRASE':'手动-词组','BROAD':'手动-广泛'}[r['匹配类型']]
    if str(r['投放']).startswith('asin'):return '商品定向-ASIN'
    if str(r['投放']).startswith('category'):return '商品定向-类目'
    if str(r['投放']).startswith('keyword-group'):return '关键词组定向'
    return '其他-待核查'
a['投放类型']=a.apply(target,axis=1)
acols=['广告曝光','广告点击','广告花费','广告销售额','广告订单','广告销售件数','广告SKU销售额','其他SKU销售额']
def adrates(d):
    d=d.copy()
    for label,x,y in [('广告CTR','广告点击','广告曝光'),('广告订单CVR','广告订单','广告点击'),
        ('广告件数每点击','广告销售件数','广告点击'),('CPC','广告花费','广告点击'),('ACOS','广告花费','广告销售额'),
        ('ROAS','广告销售额','广告花费'),('其他SKU收入占比','其他SKU销售额','广告销售额')]:
        d[label]=d[x]/d[y].replace(0,np.nan)
    return d
def adagg(keys):return adrates(a.groupby(keys,dropna=False)[acols].sum().reset_index()).sort_values('广告花费',ascending=False)
adwords=adagg(['搜索词','客户词类型']);admap=adwords.set_index('搜索词')
portfolio=adagg(['广告组合名称']);types=adagg(['投放类型'])
pq=adagg(['广告组合名称','搜索词','客户词类型'])
pq=pq.merge(a.groupby(['广告组合名称','搜索词']).agg(原稿首条行=('原稿来源行','min'),原稿末条行=('原稿来源行','max'),
    首次日期=('开始日期','min'),末次日期=('结束日期','max')).reset_index(),on=['广告组合名称','搜索词'],validate='one_to_one')
pq['队列']=np.select([pq['广告订单'].eq(0)&pq['广告点击'].ge(20)&pq['广告花费'].ge(20),
    pq['广告订单'].ge(5)&pq['广告花费'].ge(50)&pq['ACOS'].gt(.4),
    pq['广告订单'].ge(10)&pq['ACOS'].le(.25)],
    ['零单待复核','有单高成本待复核','增长测试候选'],default='观察')
candidates=pq[pq['队列'].ne('观察')].sort_values('广告花费',ascending=False)
targets=adagg(['广告组合名称','广告活动名称','广告组名称','投放','匹配类型','投放类型'])

headers=[
 '搜索词','收录ASIN数','各ASIN查询排名范围','市场搜索查询量','市场曝光','样本ASIN曝光合计',
 '平均每查询市场曝光','平均每查询样本曝光','样本曝光份额','市场点击','市场查询点击率','市场曝光CTR',
 '样本ASIN点击合计','样本曝光CTR','样本点击份额','市场点击价格中位数','点击代表ASIN价格中位数',
 '市场点击_当日达','市场点击_1天达','市场点击_2天达','市场加购','市场查询加购率','市场曝光加购比',
 '市场点击后加购比','样本ASIN加购合计','样本查询加购比','样本曝光加购比','样本点击后加购比',
 '市场加购价格中位数','加购代表ASIN价格中位数','市场加购_当日达','市场加购_1天达','市场加购_2天达',
 '市场购买','市场查询购买率','市场曝光购买比','市场点击后购买比','样本ASIN购买合计','样本购买份额',
 '样本查询购买比','样本曝光购买比','样本点击后购买比','市场购买价格中位数','购买代表ASIN价格中位数',
 '市场购买_当日达','市场购买_1天达','市场购买_2天达','报告日期','点击吸引力指数','购买承接指数',
 '广告曝光','广告点击','广告花费','广告销售额','广告销售件数','广告CTR','广告订单CVR','广告订单',
 'CPC','ACOS','ROAS','其他SKU收入占比','点击代表ASIN','加购代表ASIN','购买代表ASIN','价格说明',
 '两表收录状态','点击样本条件','初步诊断','口径提醒','样本加购份额']
assert len(headers)==71
# One-based Excel columns are retained through BE; natural differences are replaced by orders/cost metrics.
ratio_defs={7:(5,4),8:(6,4),9:(6,5),11:(10,4),12:(10,5),14:(13,6),15:(13,10),
    22:(21,4),23:(21,5),24:(21,10),26:(25,4),27:(25,6),28:(25,13),
    35:(34,4),36:(34,5),37:(34,10),39:(38,34),40:(38,4),41:(38,6),42:(38,13),
    49:(14,12),50:(42,37),56:(52,51),57:(58,52),59:(53,52),60:(53,54),61:(54,53),71:(25,21)}
percent_cols={9,11,12,14,15,22,23,24,26,27,28,35,36,37,39,40,41,42,56,57,60,62,71}
price_cols=['点击量：价格(中位数)','点击量：ASIN 售价（中位数）','加入购物车：售价（中位数）',
    '加入购物车：ASIN 价格（中位数）','购买次数：价格（中位数）','购买：ASIN 价格（中位数）']
def priceval(v):return float(v) if pd.notna(v) and 0<float(v)<=1000 else None
price_bad=s[s[price_cols].gt(1000).any(axis=1)].copy()
def representative(g,count,price):
    if g[count].max()<=0:return None,None
    r=g.sort_values([count,'ASIN'],ascending=[False,True]).iloc[0]
    return r['ASIN'],priceval(r[price])
rows=[]
for q in original_order:
    g=groups[q];m=first.loc[q];v=[None]*len(headers)
    def put(col,val):v[col-1]=clean(val)
    put(1,q);put(2,g['ASIN'].nunique())
    ranks=g['搜索查询得分'];put(3,str(int(ranks.min())) if ranks.min()==ranks.max() else f'{int(ranks.min())}—{int(ranks.max())}')
    for col,field in [(4,'搜索查询量'),(5,'曝光：曝光总量'),(10,'点击量：总次数'),(21,'购物车添加：总数'),(34,'购买：下单总数'),
        (18,'点击量：当日达配送速度'),(19,'点击量：1 天配送速度'),(20,'点击：2天配送速度'),
        (31,'购物车添加：当日达配送速度'),(32,'加入购物车：1 天配送速度'),(33,'购物车添加：2 天配送速度'),
        (45,'下单成交：当日达配送速度'),(46,'下单成交：1 天配送速度'),(47,'下单：2天配送速度'),(48,'报告日期')]:put(col,m[field])
    for col,field in zip([6,13,25,38],owncols):put(col,g[field].sum())
    for col,field in [(16,price_cols[0]),(29,price_cols[2]),(43,price_cols[4])]:put(col,priceval(m[field]))
    reps=[representative(g,field,price) for field,price in zip(owncols[1:],price_cols[1::2])]
    for (asin,price),pc,ac in zip(reps,[17,30,44],[63,64,65]):put(pc,price);put(ac,asin)
    pnotes=[]
    for label,field in zip(['点击','加购','购买'],price_cols[1::2]):
        if g[field].dropna().nunique()>1:pnotes.append(label+'价格不同；保留代表ASIN及明细')
    if g[price_cols].gt(1000).any().any():pnotes.append('原始价格>1000已标记，汇总价留空')
    put(66,'；'.join(pnotes) if pnotes else '代表价为指定ASIN中位数，不是合并中位数')
    if q in admap.index:
        ar=admap.loc[q]
        for col,field in [(51,'广告曝光'),(52,'广告点击'),(53,'广告花费'),(54,'广告销售额'),(55,'广告销售件数'),(58,'广告订单'),(62,'其他SKU收入占比')]:put(col,ar[field])
        put(67,'两表都收录')
    else:put(67,'仅SQP收录；广告未匹配')
    for col,(num,den) in ratio_defs.items():put(col,div(v[num-1],v[den-1]) if v[num-1] is not None and v[den-1] is not None else None)
    # Retain counts and absolute rates for every query, but hide unsupported multipliers.
    if v[12]<INDEX_MIN_CLICKS:
        put(49,None);put(50,None)
    put(68,'≥100点击，仍需核查稳定性' if v[12]>=100 else '<100点击，先观察')
    if v[12]<100 or v[48] is None or v[49] is None:diag='样本不足或基准无定义'
    elif v[48]>=1 and v[49]>=1:diag='吸引与承接达基准；结合相关性和成本测试'
    elif v[48]>=1 and v[49]<1:diag='能吸引点击；优先查购买承接与词意图'
    elif v[48]<1 and v[49]>=1:diag='承接达基准；优先查点击吸引环节'
    else:diag='吸引与承接均偏低；先查相关性'
    put(69,diag);put(70,'样本非全量品牌；广告范围未对齐；不相减求自然流量')
    rows.append(v)
total=pd.DataFrame(rows,columns=headers)
look=total.set_index('搜索词')
S,A=set(groups),set(admap.index)
coverage=[]
for label,terms in [('两表都收录',S&A),('仅SQP收录',S-A),('仅广告收录',A-S)]:
    d=a[a['搜索词'].isin(terms)]
    coverage.append({'集合':label,'词数':len(terms),'SQP样本点击':s[s['搜索词'].isin(terms)]['点击量：ASIN 数量'].sum() if label!='仅广告收录' else None,
        '广告点击':d['广告点击'].sum() if label!='仅SQP收录' else None,
        '广告花费':d['广告花费'].sum() if label!='仅SQP收录' else None,
        '广告销售额':d['广告销售额'].sum() if label!='仅SQP收录' else None,
        '广告订单':d['广告订单'].sum() if label!='仅SQP收录' else None})
coverage=pd.DataFrame(coverage)
adonly=adwords[adwords['搜索词'].isin(A-S)].copy()
adonly['边界']='原SQP词总表未覆盖；ASIN编码与文本分开看，不推定全是长尾'
adonly_focus=adonly[(adonly['广告点击']>=20)|(adonly['广告花费']>=20)|(adonly['广告订单']>=5)]
details=s[['ASIN','品名','颜色','规格','搜索查询','搜索查询得分','曝光量：ASIN 计数','点击量：ASIN 数量','加购：ASIN 数量','下单成交：ASIN 数量']+price_cols+['原稿来源行']].copy()
details['样本ASIN曝光CTR']=details['点击量：ASIN 数量']/details['曝光量：ASIN 计数'].replace(0,np.nan)
details['样本ASIN点击后加购比']=details['加购：ASIN 数量']/details['点击量：ASIN 数量'].replace(0,np.nan)
details['样本ASIN点击后购买比']=details['下单成交：ASIN 数量']/details['点击量：ASIN 数量'].replace(0,np.nan)
details['价格异常']=np.where(details[price_cols].gt(1000).any(axis=1),'异常待核查','')
s['同词市场率预期购买']=s['点击量：ASIN 数量']*s['购买：下单总数']/s['点击量：总次数'].replace(0,np.nan)
products=s.groupby(['ASIN','品名','颜色','规格'],dropna=False).agg(收录词数=('搜索词','size'),样本曝光=('曝光量：ASIN 计数','sum'),
    样本点击=('点击量：ASIN 数量','sum'),样本加购=('加购：ASIN 数量','sum'),样本购买=('下单成交：ASIN 数量','sum'),同词市场率预期购买=('同词市场率预期购买','sum')).reset_index()
products['样本曝光CTR']=products['样本点击']/products['样本曝光'].replace(0,np.nan)
products['样本点击后购买比']=products['样本购买']/products['样本点击'].replace(0,np.nan)
products['结构调整购买指数']=products['样本购买']/products['同词市场率预期购买'].replace(0,np.nan)
products=products.sort_values('样本点击',ascending=False)

audit=pd.DataFrame([
 ['可保留','原始SQP合并、广告词搜索数据','3,117条SQP和61,169条广告与项目原始数据逐值一致','在原稿副本上完善，原稿不变。'],
 ['可保留','总表F/M/Y/AL','样本曝光3,416,240、点击47,159、加购16,730、购买3,778，逐词求和全部正确','将ASIN标题改为已导出ASIN样本合计，避免误当单个ASIN。'],
 ['可保留','总表D/E/J/U/AH','同词市场计数一致，取一次正确','保留市场按词去重；不是把所有ASIN的市场总数相加。'],
 ['需修正','总表BC/BE/BH','BC取广告U列销售件数，BE用件数/点击，却标CVR/下单；BH相减也混合购买计数和件数','保留销售件数另列，新增广告订单T列，CVR改订单/点击。'],
 ['需修正','总表BF:BH','原稿自然曝光505行、自然点击912行、自然购买537行负值','删除自然推算，替换为订单、CPC、ACOS、ROAS等可验证指标。'],
 ['需修正','总表AB/AP/AM','分别39、39、537个除零错误，共615个单元格','分母0或缺失时留空，并在说明中区分无定义与真实0。'],
 ['需限定','总表Q/AD/AR/C','有371个词对应多个ASIN；价格和查询排名不一定一致，用None拒绝不一致方向正确','代表价明确对应各环节计数最多的ASIN，完整明细保留；查询排名显示范围而非品牌排名。'],
 ['需修正','总表Z/AN标题','查询加购比=样本加购/查询量，不等于原始ASIN加购份额；查询购买比同理','分别命名；补样本加购份额=样本加购/市场加购。'],
 ['需补充','原稿总表仅2,410个SQP词','仅匹配56,382广告点击，占全表35.56%；未覆盖42,044个广告客户条目','新增全量广告概览、覆盖表、仅广告重点和完整CSV。'],
 ['需补充','原稿BA/BB','有花费和销售额但没有ACOS、CPC等成本指标','新增成本及SKU收入构成；目标ACOS和利润仍待补。'],
 ['需优化','总表LET/FILTER与SUMIF整列引用','大量整列计算易慢且依赖新版Excel函数','计数按版本预计算，比例用简单公式并缓存；源表更新后运行脚本刷新计数。'],
],columns=['判断','原稿位置','证据','完善方式'])
notes=pd.DataFrame([
 ['范围','2026年9月；51个SQP ASIN×词样本，原稿内两张原始数据页；美国USD广告报告。'],
 ['输入版本','reference/26.9汇总表_原稿.xlsx；原稿SHA256：'+src_hash],
 ['样本计数','总表样本ASIN曝光/点击/加购/购买是多个已导出ASIN求和，不是单个ASIN，也不是完整品牌。'],
 ['查询指标','点击、加购、购买÷搜索查询量，表示每次查询对应的事件计数，不是去重用户概率。'],
 ['曝光CTR','点击÷曝光。样本CTR分子分母均为同词已收录ASIN计数；市场CTR用同词市场计数。'],
 ['份额','样本曝光/市场曝光、样本点击/市场点击、样本加购/市场加购、样本购买/市场购买。'],
 ['自然流量','理论上同范围完整总量减付费量可拆分；现有两份报告未对齐，当前不计算自然量或广告占比。'],
 ['广告订单和件数','订单取广告源表T列，件数取U列。广告CVR用订单/点击，件数保留但不当作订单。'],
 ['代表价格','点击价取点击最多的ASIN；加购价取加购最多的ASIN；购买价取购买最多的ASIN。并列按ASIN字典序选择。没有该事件时留空；均不是合并中位数。'],
 ['价格异常','原始价格>1000的9行标记待核查；总表不使用异常代表价，原始数据和明细保留数值。'],
 ['零与缺失','已记录的0计数保留0；没匹配到广告或分母0的比例留空，不填0。'],
 ['指数','吸引力=样本CTR/市场CTR；承接=样本点击后购买比/市场点击后购买比。样本点击<100或基准无定义时留空。100点击只是展示筛选线，不是显著性标准；达到后仍须看曝光、点击、购买计数和绝对比率。低市场基准会放大倍数，不按指数大小判断机会优先级。'],
 ['极小样本案例','hal：样本1曝光1点击，CTR100%；市场1,648点击/10,764,360曝光，CTR约0.01531%。原始比值约6,532倍，不能据此判断吸引力强；完善版保留计数和CTR，两个指数留空。'],
 ['结构调整','实际购买/Σ每词样本点击×同词市场购买/点击，控制已观察词结构，不控制价格、位置、人群或缺失词。'],
 ['筛选线','零单≥20点击且花费≥20；高成本≥5单、花费≥50且ACOS>40%；增长≥10单且ACOS≤25%。均为待复核规则，不是目标或盈亏线。'],
 ['复算','计数和分组是本次输入版本的计算快照；总表比例为Excel公式且有缓存。原始页修改后，运行review_user_summary.py刷新快照。'],
 ['日期','11,646行广告覆盖跨日期区间，占71.6%花费，不以开始日期生成每日趋势。'],
 ['追溯','ASIN明细行号对应原稿原始SQP合并页；组合词行号对应原稿广告词搜索数据页。首末行号只定位范围。'],
 ['阅读顺序','先看结论，再看重点词看板和行动候选，需要全字段时看总表，再通过原稿行号查源数据。'],
],columns=['主题','说明'])

totals=adrates(pd.DataFrame([a[acols].sum()]))
stats={'original_sha256':src_hash,'input_sqp_rows':len(s),'input_ad_rows':len(a),'sqp_terms':len(S),'ad_terms':len(A),
    'error_cells':dict(errors),'negative_cells':dict(negative),'multi_asin_terms':int(s.groupby('搜索词')['ASIN'].nunique().gt(1).sum()),
    'own':{col:int(s[col].sum()) for col in owncols},'ad_totals':totals.iloc[0].to_dict(),'coverage':coverage.to_dict(orient='records'),
    'sample_purchase_per_click':div(s['下单成交：ASIN 数量'].sum(),s['点击量：ASIN 数量'].sum()),
    'market_purchase_per_click':div(first['购买：下单总数'].sum(),first['点击量：总次数'].sum()),
    'own_structure_market_purchase_per_click':div(s['同词市场率预期购买'].sum(),s['点击量：ASIN 数量'].sum())}
def safe_json(v):
    if isinstance(v,dict):return {k:safe_json(x) for k,x in v.items()}
    if isinstance(v,list):return [safe_json(x) for x in v]
    return clean(v)
(OUT/'review_metrics.json').write_text(json.dumps(safe_json(stats),ensure_ascii=False,indent=2,default=str,allow_nan=False))
for name,d in [('校验与修改记录',audit),('完善总表',total),('广告覆盖',coverage),('广告客户词全量',adwords),
    ('广告仅出现全量',adonly),('广告组合词全量',pq),('行动候选',candidates),('ASIN明细',details),('商品汇总',products)]:
    d.to_csv(OUT/(name+'.csv'),index=False,encoding='utf-8-sig')
print(json.dumps(safe_json(stats),ensure_ascii=False,indent=2,default=str))

def money(x):return f'${x:,.2f}'
def pct(x):return f'{x:.1%}'
def mdtable(d,cols,fmts=None):
    fmts=fmts or {}
    lines=['| '+' | '.join(cols)+' |','| '+' | '.join(['---']*len(cols))+' |']
    for _,r in d.iterrows():
        vals=[]
        for col in cols:
            val=r[col]
            if pd.isna(val):val='—'
            elif col in fmts:val=fmts[col](val)
            elif isinstance(val,(float,np.floating)):val=f'{val:,.0f}' if val==int(val) else f'{val:.2f}'
            vals.append(str(val).replace('|','/'))
        lines.append('| '+' | '.join(vals)+' |')
    return '\n'.join(lines)
fmt={'样本曝光CTR':pct,'样本点击后购买比':pct,'结构调整购买指数':lambda x:f'{x:.2f}倍',
    '样本ASIN点击合计':lambda x:f'{x:,.0f}','样本ASIN购买合计':lambda x:f'{x:,.0f}',
    '市场点击后购买比':pct,'广告花费':money,'广告销售额':money,'ACOS':pct,'其他SKU收入占比':pct}
top=products.iloc[0]
evidence=pd.DataFrame([
 ['已核对事实','原稿的核心求和正确','逐词核对曝光、点击、加购、购买，无差异；对应样本合计3,416,240、47,159、16,730、3,778','保留已有市场按词去重、样本ASIN计数求和的思路。'],
 ['已核对事实','总表不覆盖全部广告','总表2,410个SQP词，仅匹配56,382广告点击，占全表35.56%；还有42,044个仅广告客户条目','看全量广告概览和覆盖页，不能用总表广告合计代表全店广告。'],
 ['已核对事实','需要区分广告订单和件数','16,016个广告订单、16,571件；订单CVR10.10%，件数/点击10.45%','用订单/点击作广告CVR，保留件数另列。'],
 ['诊断线索','先查主力商品承接','64oz黑色单装B0DGCVK1JL占样本点击70.5%、购买61.9%，结构调整购买指数0.77','检查大词意图、规格、页面、价格优惠、库存、评价与展示位置；原因未确定。'],
 ['诊断线索','mason jars吸引和承接均偏低','1,549点击60购买；样本CTR0.56% vs市场1.29%，购买/点击3.87% vs12.42%','先确认泛储存罐意图与带把手壶型是否适配，再查页面和投放。'],
 ['诊断线索','cold brew pitcher主要查承接','1,181点击69购买；CTR2.48% vs市场2.24%，购买/点击5.84% vs11.97%','检查冷萃功能与词意图、页面、价格和流量构成；不能直接说页面有问题。'],
 ['已核对事实','同词组合广告效率不同','mason jar pitcher在梅森罐组合ACOS22.9%、437单，32oz组合47.3%、91单','定位到具体活动/广告组/推广商品，分别判断，避免整店统一否词或加价。'],
 ['待验证建议','明确场景词增长测试','梅森罐组合sun tea glass jar for outside：323点击84单、ACOS20.3%；tea pitcher for fridge：234点击60单、19.6%','确认利润、库存和预算约束，先小额单变量测试。'],
 ['待验证建议','高成本复核队列','45个组合词有单且ACOS>40%、花费≥50，合计花费6,945.61；不是可节省金额','先查具体投放单元与利润，不直接暂停所有候选。'],
 ['待验证建议','零单候选规模有限','按≥20点击且花费≥20筛出26个组合词，花费702.04，仅占全部花费0.57%','查相关性、归因成熟度与具体目标；20点击零单仍可能随机发生。'],
 ['已核对事实','总ACOS需结合销售构成','总ACOS29.73%，其他SKU归因销售额占44.01%','没有商品成本和实际售出SKU，不能判断利润或广告增量。'],
 ['当前不能确定','自然量、净利润与差异原因','ASIN/广告场景范围和购买归因未对齐；无成本、页面、库存等直接证据','不将脱敏、价格或页面问题定为原因；不相减推自然流量。'],
],columns=['证据层级','发现','数据证据','下一步或边界'])

casewords=['mason jars','cold brew pitcher','mason jar pitcher','glass pitcher with lid','sun tea glass jar for outside','tea pitcher for fridge','glass carafe with lid']
cases=total[total['搜索词'].isin(casewords)].copy()
ad_case=pq[(pq['搜索词'].eq('mason jar pitcher'))&pq['广告花费'].ge(100)].sort_values('ACOS')
brief='''# 26.9汇总表：完善说明与分析结论

## 这份表现在能用到什么程度

这份表可以作为词级分析基础。内嵌3,117条SQP和61,169条广告与项目原始数据一致；市场同词计数一致后取一次、ASIN计数按词求和的思路成立，四项样本求和逐词核对全部正确。你已经把查询点击率、曝光CTR、查询加购率和点击后加购比分开，后续可以沿这个结构分析。

完善版保留你的核心字段顺序，并添加易读的重点词看板、证据层级、广告成本、覆盖范围、ASIN明细和行动候选。原稿保存在reference/26.9汇总表_原稿.xlsx，未修改。

## 1. 需要完善的地方

'''+mdtable(audit,['判断','原稿位置','证据','完善方式'])+'''

三个容易误解的细节：

- 总表中的ASIN点击、曝光和购买是多个已导出ASIN的样本合计，不能称为某一个ASIN的CTR或完整品牌总量。
- 价格和查询排名不一致时显示None，避免了任意取某个ASIN值。完善版用各环节计数最多的ASIN作为代表，明确其ASIN和中位数；不是合并中位数。不同ASIN的真实合并中位数不能从中位数和事件数量精确还原。
- 分母为0时指标无定义，留空；真正有分母但分子为0时可以显示0。没匹配到广告也留空，不等于该词没有任何广告。

## 2. 表中广告数据只覆盖一部分

'''+mdtable(coverage,['集合','词数','SQP样本点击','广告点击','广告花费','广告销售额','广告订单'],fmt)+'''

总表匹配的广告花费$52,930.53，完整广告花费$122,725.96。仅广告收录条目贡献102,186点击和$69,795.43花费；其中含商品ASIN编码，不全部是普通长尾查询。看广告整体效率和降本，必须同时看这些条目。ASIN编码的出现不单独证明广告位置。

## 3. 可以先讲清楚的分析结论

### 主力商品的检查优先级最高

'''+mdtable(products.head(6),['ASIN','品名','样本点击','样本购买','样本点击后购买比','结构调整购买指数'],fmt)+'''

64oz黑色单装B0DGCVK1JL贡献33,260点击和2,339购买，占样本70.5%的点击、61.9%的购买计数。优先检查它是因为体量大、影响面广，不代表已经证明它有页面问题。

“结构调整购买指数”=实际购买÷Σ每词样本点击×同词市场购买/点击。B0DGCVK1JL为0.77，表示按已观察点击词结构与同词市场率对照，实际购买计数约为基准的77%。指标用于诊断排序，不是因果或显著性检验，也不代表能追回多少订单。

### 词级对照比只看总率更有用

'''+mdtable(cases,['搜索词','样本ASIN点击合计','样本ASIN购买合计','样本曝光CTR','市场曝光CTR','样本点击后购买比','市场点击后购买比'],
    {**fmt,'样本曝光CTR':pct,'市场曝光CTR':pct})+'''

mason jars样本CTR0.56%低于市场1.29%，购买/点击3.87%低于12.42%，先查用户想要普通储存罐还是带把手倒水壶，以及实际商品是否适配。cold brew pitcher样本CTR2.48%略高于市场2.24%，但购买/点击5.84%低于11.97%，优先检查购买承接与功能意图。轻微率差不能自动当显著差异。

点击率只能用于同词同口径比较；不能把曝光、点击、购买差异直接定为自然或广告造成。

### 汇总率可能掩盖词结构

样本整体购买/点击8.01%，覆盖词市场总体为7.69%。直接看这两个数，容易以为样本整体优于市场；但样本点击集中在不同的词上。按样本自身的点击词结构加权，同词市场基准为9.41%，样本实际仍为8.01%，指数0.85。

这些对照分别回答不同问题。当前数据支持进一步检查承接，不能据此断言商品页面有问题，也不能将与基准的差额当成可挽回订单。

### 同一个词不能全店统一调整

'''+mdtable(ad_case,['广告组合名称','搜索词','广告点击','广告订单','广告花费','ACOS'],fmt)+'''

同词mason jar pitcher在【梅森罐】为22.9%，在【32oz梅森罐】为47.3%。应拆到具体活动、广告组、投放目标和推广商品再判断。广告组合名不是可靠的ASIN映射，不能将差异直接归因于容量或颜色。

### 使用场景是一个可验证的增长方向

【梅森罐】组合的sun tea glass jar for outside为323点击84单、ACOS20.3%；tea pitcher for fridge为234点击60单、19.6%；glass carafe with lid为63点击21单、13.3%。SQP对应词也有一定承接线索，适合在确认商品实际形态、利润、库存和预算约束后测试。

对照范围仍不完全相同，不能将全组合广告表现直接归到某一个ASIN。glass carafe with lid只有145个SQP样本点击，先验证稳定性，不预言稳定增长。iced tea pitcher的SQP表现相对较好，但该词全表广告ACOS31.9%，说明搜索表现好也需要单独判断广告成本。

### 降本先查具体单元，零单总花费不是可节省预算

按“组合×客户词”筛出45个有单高成本候选，合计花费$6,945.61，均有成交，适合查成本和利润，不能直接停投。26个零单复核候选合计$702.04，仅占全部花费0.57%。其筛选线为≥20点击且花费≥$20，仍可能有随机波动；真实购买率10%时，20点击零单的概率约12%。

更有体量的复核对象包括：【梅森罐】低价扫流0.7 -2 - z的substitutes，花费$3,654.70、346单、ACOS41.9%；【32oz梅森罐】低价扫流0.7 - z的complements，$944.69、77单、ACOS53.3%。先拆具体商品目标和连带销售再测试。

候选中的25%和40%只是描述性筛选线，用户尚未提供目标ACOS、盈亏线和商品利润。

## 4. 这次不能下的结论

自然流量占比、广告增量贡献、净利润、页面/价格/库存的具体因果影响，当前都缺少必要证据。总ACOS29.73%包含44.01%的其他SKU归因收入，它有业务价值，但不等于被投商品盈利或广告新增销售。

广告报告有11,646行开始日期与结束日期不同，覆盖71.6%花费，因此不画每日趋势。不能用脱敏作为解释所有差异的确定原因；现有数据能验证的是导出范围和百分比显示精度。

## 5. 你可以怎样继续完善

先确定分析问题，再选指标：看点击吸引用曝光CTR；看点击后承接用加购/点击、购买/点击；看竞争地位用曝光、点击、加购、购买份额；看广告成本用CPC、ACOS、订单CVR和SKU销售构成。查询口径指标和曝光/点击口径指标分别保留，但不直接比较百分比大小。

两个指数仅作同词相对比率的辅助描述，不作为机会排名。样本点击少于100时，总表和看板均留空；100点击是工作筛选线，不是统计显著性或稳定性保证。市场基准很低时，即使达到筛选线，倍数仍可能很大，须一起看原始计数、绝对比率及百分点差。购买承接还须核查购买计数；不能把零购买自动判为无承接能力。

例如hal只有1次样本曝光、1次点击，CTR为100%；市场为1,648次点击/10,764,360次曝光，CTR约0.01531%，原始比值约6,532倍。这不能证明吸引力强。再出现1次曝光而没有点击，样本CTR就会降到50%。当前保留原始计数与CTR，隐藏小样本倍数，不修改原始数据或截断极端值。

每条结论按“事实→可能原因→验证动作”写。例如：cold brew pitcher有1,181点击69购买，同词购买率低于市场；可能与功能意图、价格、页面或流量位置有关；先检查各ASIN、实际投放商品和页面资料，再做小范围单变量测试。

下一步最有价值的数据是可靠的推广ASIN/SKU映射、各商品广告前贡献利润率、库存与断货记录、价格优惠、评论变化、日粒度广告报告及前后月份同口径数据。补这些数据后，再区分成本问题、商品承接问题和可增长机会。

## 6. 文件使用与更新

优先打开26.9汇总表_完善版.xlsx中的“先看结论”“重点词看板”“行动候选”。全字段在“总表”；价格和具体ASIN在“ASIN明细”；全量广告概览、广告覆盖和仅广告重点补足原稿范围。

计数和分组是输入版本的计算快照，比例为简单Excel公式且写入缓存，不依赖LET/FILTER和整列扫描。修改源数据后，在项目根目录运行python user_summary_review/review_user_summary.py更新快照；只手动修改计数时，比例会自动重新计算。原始两页的数据内容保留在完善版，原稿另存未改动。

完整广告客户词、仅广告条目和组合词明细也输出为CSV，不因看板筛选而删掉低点击条目。行号对应原稿源表，首末行号仅定位范围，不表示中间每行均属于同一词。
'''
(OUT/'26.9汇总表_完善说明与分析结论.md').write_text(brief)
boardcols=['搜索词','市场搜索查询量','收录ASIN数','样本ASIN曝光合计','样本ASIN点击合计','样本ASIN加购合计','样本ASIN购买合计',
    '样本曝光CTR','市场曝光CTR','样本点击后加购比','市场点击后加购比','样本点击后购买比','市场点击后购买比',
    '点击吸引力指数','购买承接指数','样本购买份额','广告点击','广告订单','广告花费','ACOS','CPC','其他SKU收入占比',
    '两表收录状态','点击样本条件','初步诊断']
board=total[boardcols].sort_values('样本ASIN点击合计',ascending=False)

formula_info=pd.DataFrame([{'完善版列':get_column_letter(col),'指标':headers[col-1],
    'Excel第2行公式':excel_ratio(get_column_letter(num)+'2',get_column_letter(den)+'2','M2' if col in (49,50) else None),
    '留空条件':'样本点击<100或基准无定义；100为展示筛选线，不保证稳定' if col in (49,50) else '输入非数字或分母≤0；指标无定义'} for col,(num,den) in ratio_defs.items()])

# Stream the two large source sheets instead of loading a million-cell editable workbook.
wb=openpyxl.Workbook(write_only=True)
wb.calculation=CalcProperties(fullCalcOnLoad=True,forceFullCalc=True,calcMode='auto')
dark='173D59';header_font=Font(name='微软雅黑',size=10,bold=True,color='FFFFFF')
body_font=Font(name='微软雅黑',size=10)
header_fill=PatternFill('solid',fgColor=dark)
header_align=Alignment(wrap_text=True,vertical='center')
sheet_sizes={};formula_cache={}
def add_sheet(name,frame,raw=False,formulas=False):
    ws=wb.create_sheet(name);ws.freeze_panes='C2' if len(frame.columns)>15 else 'A2'
    ws.sheet_view.showGridLines=False;ws.row_dimensions[1].height=45
    hs=list(frame.columns)
    for j,h in enumerate(hs,1):
        if raw:width=32 if isinstance(h,str) and h in ['客户搜索词','广告活动名称','广告组名称','投放','搜索查询'] else 18
        elif h in ['说明','数据证据','下一步或边界','证据','完善方式','边界','初步诊断','口径提醒','价格说明']:width=60
        elif h in ['搜索词','搜索查询','原始客户词','品名','广告活动名称','广告组名称','投放']:width=38
        else:width=max(14,min(30,len(str(h))*2+2))
        ws.column_dimensions[get_column_letter(j)].width=width
    cells=[]
    for h in hs:
        c=WriteOnlyCell(ws,str(h));c.data_type='s';c.font=header_font;c.fill=header_fill;c.alignment=header_align;cells.append(c)
    ws.append(cells)
    for rownum,record in enumerate(frame.itertuples(index=False,name=None),2):
        cells=[]
        if len(frame)<100:ws.row_dimensions[rownum].height=90 if name in ['先看结论','口径说明','校验与修改'] else 34
        for j,(h,value) in enumerate(zip(hs,record),1):
            value=clean(value)
            if formulas and j in ratio_defs:
                num,den=ratio_defs[j]
                formula=excel_ratio(get_column_letter(num)+str(rownum),get_column_letter(den)+str(rownum),'M'+str(rownum) if j in (49,50) else None)
                c=WriteOnlyCell(ws,formula);formula_cache[f'{get_column_letter(j)}{rownum}']=value
            elif raw and not isinstance(value,str):
                cells.append(value);continue
            else:
                c=WriteOnlyCell(ws,value)
                if isinstance(value,str):c.data_type='s'
            if not raw:
                c.font=body_font
                c.alignment=Alignment(vertical='top',wrap_text=len(frame)<100)
                if (formulas and j in percent_cols) or (not formulas and any(x in str(h) for x in ['CTR','CVR','份额','占比','点击后购买比','点击后加购比'])):
                    c.number_format='0.00%'
                elif any(x in str(h) for x in ['价格','售价','花费','销售额','CPC']):c.number_format='"$"#,##0.00'
                elif any(x in str(h) for x in ['指数','ROAS','每查询','件数每点击','预期购买']):c.number_format='0.00'
                elif isinstance(value,(int,float)):c.number_format='#,##0'
            cells.append(c)
        ws.append(cells)
    ws.auto_filter.ref=f'A1:{get_column_letter(len(hs))}{len(frame)+1}'
    sheet_sizes[name]=(len(frame)+1,len(hs))
    return ws

add_sheet('先看结论',evidence)
add_sheet('重点词看板',board)
add_sheet('总表',total,formulas=True)
add_sheet('口径说明',notes)
add_sheet('公式说明',formula_info)
add_sheet('校验与修改',audit)
add_sheet('ASIN明细',details)
add_sheet('商品汇总',products)
add_sheet('广告全量概览',totals)
add_sheet('广告覆盖',coverage)
add_sheet('广告组合',portfolio)
add_sheet('投放类型',types)
add_sheet('行动候选',candidates)
add_sheet('投放单元',targets)
add_sheet('仅广告重点',adonly_focus)
add_sheet('价格异常',price_bad[['ASIN','搜索查询']+price_cols+['原稿来源行']])
add_sheet('原始SQP合并',sraw,raw=True)
add_sheet('广告词搜索数据',araw,raw=True)
intermediate=OUT/'_workbook_uncached.xlsx';wb.save(intermediate)

# Cache only the simple ratio formulas, retain them for normal Excel editing, and set dimensions.
NS='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
ET.register_namespace('',NS)
with zipfile.ZipFile(intermediate) as zin:
    sheets=ET.fromstring(zin.read('xl/workbook.xml')).find('{'+NS+'}sheets')
    rels={r.attrib['Id']:r.attrib['Target'] for r in ET.fromstring(zin.read('xl/_rels/workbook.xml.rels'))}
    paths={}
    for sh in sheets:
        target=rels[sh.attrib['{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id']]
        path=target.lstrip('/') if target.startswith('/') else 'xl/'+target
        paths[path]=sh.attrib['name']
    with zipfile.ZipFile(BOOK,'w',zipfile.ZIP_DEFLATED,compresslevel=6) as zout:
        for item in zin.infolist():
            content=zin.read(item.filename)
            if item.filename in paths:
                name=paths[item.filename]
                if name=='总表':
                    xml=ET.fromstring(content)
                    for c in xml.findall('.//{'+NS+'}sheetData/{'+NS+'}row/{'+NS+'}c'):
                        addr=c.attrib['r']
                        if addr in formula_cache:
                            val=formula_cache[addr];node=c.find('{'+NS+'}v')
                            if node is None:node=ET.SubElement(c,'{'+NS+'}v')
                            node.text=str(val) if val is not None else None
                    nr,nc=sheet_sizes[name]
                    dim=ET.Element('{'+NS+'}dimension',{'ref':f'A1:{get_column_letter(nc)}{nr}'})
                    xml.insert(1,dim);content=ET.tostring(xml,encoding='utf-8')
                else:
                    nr,nc=sheet_sizes[name]
                    # Generated write-only worksheets lack a dimension; inserting it avoids scanning on open.
                    marker=b'</sheetPr>'
                    dim=f'<dimension ref="A1:{get_column_letter(nc)}{nr}"/>'.encode()
                    assert marker in content
                    content=content.replace(marker,marker+dim,1)
            zout.writestr(item,content)
intermediate.unlink()
assert hashlib.sha256(INPUT.read_bytes()).hexdigest()==src_hash

# Standalone Word copy of the findings for reading and sharing.
from docx import Document
from docx.shared import Inches,Pt
from docx.oxml.ns import qn
doc=Document();section=doc.sections[0]
section.left_margin=Inches(.7);section.right_margin=Inches(.7)
normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(10)
normal._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'微软雅黑')
def text_only(s):return s.replace('**','').replace('`','')
lines=brief.splitlines();i=0
while i<len(lines):
    line=lines[i].strip()
    if not line:i+=1;continue
    if line.startswith('|'):
        block=[]
        while i<len(lines) and lines[i].strip().startswith('|'):
            block.append([text_only(x.strip()) for x in lines[i].strip().strip('|').split('|')]);i+=1
        tbl=doc.add_table(rows=0,cols=len(block[0]));tbl.style='Light Shading Accent 1'
        for r in [block[0]]+block[2:]:
            cells=tbl.add_row().cells
            for cell,value in zip(cells,r):cell.text=value
        continue
    if line.startswith('# '):doc.add_heading(text_only(line[2:]),0)
    elif line.startswith('## '):doc.add_heading(text_only(line[3:]),1)
    elif line.startswith('### '):doc.add_heading(text_only(line[4:]),2)
    elif line.startswith('- '):doc.add_paragraph(text_only(line[2:]),style='List Bullet')
    else:doc.add_paragraph(text_only(line))
    i+=1
doc.save(OUT/'26.9汇总表_完善说明与分析结论.docx')

# Reading order and formulas are visible in the delivered workbook and repository.
(OUT/'README.md').write_text('''# 用户汇总表的完善版

输入：reference/26.9汇总表_原稿.xlsx，用户上传原稿，未修改。计算仅使用其两张源数据页和商品对应关系。

先打开26.9汇总表_完善版.xlsx：先看结论 → 重点词看板 → 行动候选 → 总表。共18个页签，包括完整源数据、广告范围、价格明细、校验与公式说明。

26.9汇总表_完善说明与分析结论.md / .docx解释修正、数据证据和下一步。CSV保留全部广告词、仅广告条目和组合词，不因看板筛选而丢掉低点击记录。

字段计算公式说明.xlsx / .md逐列列出看板25列与总表71列的计算、汇总、取数或不计算规则。跨表按搜索词匹配。

两个指数在样本点击<100或基准无定义时留空。100点击仅是展示筛选线，不是显著性检验；达到后仍需核查曝光、点击、购买计数和绝对比率。市场基准很低会放大倍数，不以指数大小排列机会；hal的一次点击不能支持吸引力判断。

总表的比例使用简单Excel公式并写入缓存；分组计数是输入版本快照。修改原始页后运行脚本刷新计数；只修改计数时比例正常重算。

项目根目录复算：python user_summary_review/review_user_summary.py。依赖见根目录requirements.txt。

更新字段公式指南：python user_summary_review/build_field_formula_guide.py。若直接在Excel汇总广告源表，先按指南统一客户词空格。

自然量暂不计算；样本不是完整品牌；价格为各环节代表ASIN中位数；广告订单与件数分开；筛选线不是目标ACOS；诊断不是因果结论。
''')
print('Created:',BOOK,'with',len(sheet_sizes),'sheets and',len(formula_cache),'cached ratio formulas.')
